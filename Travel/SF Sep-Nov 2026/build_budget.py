# -*- coding: utf-8 -*-
"""Build the San Francisco Sep 28 - Nov 4 2026 daily budget workbook (uploaded to Google Sheets)."""
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

OUT = "/Users/germantatis/Desktop/GTMCoding/Travel/SF Sep-Nov 2026/Presupuesto SF Sep28-Nov4 2026.xlsx"

START, END = date(2026, 9, 28), date(2026, 11, 4)
JOAQUIN_FROM = date(2026, 10, 25)   # supuesto: ultimas 10 noches en casa de Joaquin
DAYS = [START + timedelta(d) for d in range((END - START).days + 1)]
N = len(DAYS)                       # 38
FIRST, LAST = 5, 5 + N - 1          # filas 5..42 en 'Dia a dia'
DIA = "'Dia a dia'"

NAVY = "1B1F3B"; TEAL = "0F766E"; GREY = "F3F4F6"; LIGHT = "E8F1FF"; YEL = "FFF7D6"; GREEN = "E6F4EA"; RED = "FDE8E8"
hdr_font = Font(bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor=NAVY)
sub_fill = PatternFill("solid", fgColor=GREY); in_fill = PatternFill("solid", fgColor=YEL)
tot_fill = PatternFill("solid", fgColor=LIGHT)
thin = Side(style="thin", color="D0D5DD"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
USD = '"$"#,##0.00'; USD0 = '"$"#,##0'; PCT = '0.00%'; DATEF = 'dd-mmm-yyyy'
DIAS_ES = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]

def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    if sub:
        ws["A2"] = sub; ws["A2"].font = Font(italic=True, color="555555")

def header(ws, row, cols, start_col=1):
    for i, c in enumerate(cols):
        cell = ws.cell(row=row, column=start_col + i, value=c)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.border = box
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x

wb = Workbook()

# ------------------------------------------------------------------ SUPUESTOS
sp = wb.active; sp.title = "Supuestos"
title(sp, "Supuestos y escenarios", "Celdas amarillas = editables. Todo lo demás se recalcula. Montos en USD.")
sp["A3"] = "Escenario activo (mueve la tabla Día a día)"; sp["A3"].font = Font(bold=True)
sp["B3"] = "Realista"; sp["B3"].fill = in_fill; sp["B3"].font = Font(bold=True)
dv = DataValidation(type="list", formula1='"Política,Realista,Ahorro"', allow_blank=False); sp.add_data_validation(dv); dv.add("B3")
sp["A4"] = "Índice de escenario (auto)"; sp["B4"] = '=IF(B3="Política",1,IF(B3="Realista",2,3))'
sp["A5"] = "Tu parte del Airbnb (compartido con Samuel)"; sp["B5"] = 0.5; sp["B5"].number_format = PCT; sp["B5"].fill = in_fill

header(sp, 7, ["Categoría (por día / por noche)", "Política (tope Yuno)", "Realista (mercado SF)", "Ahorro", "Valor activo", "Unidad", "Fuente / nota"])
rows = [
    # row 8
    ("Airbnb: precio de la unidad por noche (2 pax, sin impuestos)", 440, 130, 121, "USD/noche",
     "Política: 2 × $220 (tope hotel SF por persona). Realista: listing SOMA Executive Retreat (2BR) aprobado por Sean como '<$4k/mes' = ~$130/noche. CONFIRMAR con el precio real del listing. Ahorro: cuartil inferior Airbnb SF $121 (AirROI 2026)."),
    # row 9
    ("Airbnb: tu parte por noche (sin impuestos)", "=B8*$B$5", "=C8*$B$5", "=D8*$B$5", "USD/noche", "Unidad × tu parte (Supuestos B5)."),
    # row 10
    ("Impuestos Airbnb (TOT 14% + TID 2,25%) si la reserva es < 30 noches", "=IF(Resumen!$B$9<30,0.1625,0)", "=IF(Resumen!$B$9<30,0.1625,0)", "=IF(Resumen!$B$9<30,0.1625,0)", "%",
     "SF Transient Occupancy Tax 14% + Tourism Improvement District 2,25% aplican solo a estadías < 30 noches (sftreasurer.org). Con 30+ noches queda en 0%."),
    # row 11
    ("Desayuno", 15, 12, 5, "USD/día", "Política: tope $15 (recibo auto <$30). Realista: café $5,50 + bagel/pastry ~$6 con tax y tip. Ahorro: mercado (Trader Joe's) en el Airbnb."),
    # row 12
    ("Almuerzo", 25, 25, 15, "USD/día", "Política: tope $25. Realista: casual $15–25 en menú + ~30% (tax 8,625% + SF Mandate 3–7% + tip 18–20%). Ahorro: fast-casual / bowl ~$12 + tax."),
    # row 13
    ("Cena", 50, 45, 20, "USD/día", "Política: tope $50 (recibo obligatorio, tip incluido). Realista: plato $20–40 + bebida + 30%. Ahorro: cocinar (groceries $400–500/mes = ~$15/día)."),
    # row 14
    ("Transporte local", 55, 15, 6, "USD/día", "Política: tope $55/día (Uber default, transporte público primero). Realista: Muni pass 'A' $104/mes + 3–4 Uber/semana × ~$25. Ahorro: solo Muni/BART ($2,85/viaje, pass M $86)."),
    # row 15
    ("Otros diarios (snacks, agua, propinas varias, farmacia)", 0, 5, 2, "USD/día", "No es una categoría de la política: sale de tu bolsillo salvo que tenga propósito de negocio."),
]
for i, (a, b, c, d, u, note) in enumerate(rows):
    r = 8 + i
    sp.cell(r, 1, a); sp.cell(r, 2, b); sp.cell(r, 3, c); sp.cell(r, 4, d)
    sp.cell(r, 5, f"=CHOOSE($B$4,B{r},C{r},D{r})"); sp.cell(r, 6, u); sp.cell(r, 7, note)
    for col in (2, 3, 4, 5):
        sp.cell(r, col).number_format = PCT if u == "%" else USD
        if col in (2, 3, 4) and r != 9 and r != 10: sp.cell(r, col).fill = in_fill
    sp.cell(r, 7).alignment = Alignment(wrap_text=True, vertical="top")

header(sp, 17, ["Gastos únicos (una sola vez en el viaje)", "Política (tope Yuno)", "Realista", "Ahorro", "Valor activo", "Unidad", "Fuente / nota"])
ones = [
    # row 18
    ("Vuelo BOG–SFO ida y vuelta (YA COMPRADO)", 1064, 1064, 1064, "USD", "Pagado. Avianca Flex vía SAL (hilo #travel-requests 14-sep). Tiquete flex: puedes volver antes."),
    # row 19
    ("Traslado SFO → Airbnb (28-sep, llegada)", 0, 45, 11.15, "USD", "Política: cabe dentro del tope de $55/día, por eso 0 extra. Realista: UberX SFO–downtown $35–55 (+$5,50 airport surcharge). Ahorro: BART $11,15."),
    # row 20
    ("Traslado casa → SFO (4-nov, salida)", 0, 11.15, 11.15, "USD", "BART desde el centro $11,15 (una vía). Sube a 45 si vas en Uber con maletas."),
    # row 21
    ("eSIM / datos móviles (38 días)", 75, 75, 57, "USD", "Reembolsable con recibo (Policy §9). T-Mobile US Pass $50/30 días + $25/7 días (se activa dentro de EE.UU.). Ahorro: Airalo 50 GB/30 días $42 + top-up."),
    # row 22
    ("Lavandería / varios únicos", 0, 20, 20, "USD", "El listing tiene lavadora. Wash & fold en SF $2,39–2,50/lb si la necesitas."),
]
for i, (a, b, c, d, u, note) in enumerate(ones):
    r = 18 + i
    sp.cell(r, 1, a); sp.cell(r, 2, b); sp.cell(r, 3, c); sp.cell(r, 4, d)
    sp.cell(r, 5, f"=CHOOSE($B$4,B{r},C{r},D{r})"); sp.cell(r, 6, u); sp.cell(r, 7, note)
    for col in (2, 3, 4, 5):
        sp.cell(r, col).number_format = USD
        if col in (2, 3, 4): sp.cell(r, col).fill = in_fill
    sp.cell(r, 7).alignment = Alignment(wrap_text=True, vertical="top")

sp["A24"] = "Comidas con clientes / equipo (número estimado en el viaje)"; sp["B24"] = 0; sp["B24"].fill = in_fill
sp["C24"] = "Tope $85 por persona (SF). Requiere etiquetar asistentes en RAMP. No cuenta contra tus topes individuales. Déjalo en 0 si no lo quieres en el presupuesto."
sp["A25"] = "Costo por comida con clientes (tope SF por persona)"; sp["B25"] = 85; sp["B25"].number_format = USD
sp["A26"] = "Total comidas con clientes / equipo"; sp["B26"] = "=B24*B25"; sp["B26"].number_format = USD; sp["B26"].font = Font(bold=True)
for r in range(1, 27): sp.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
widths(sp, [52, 18, 18, 14, 14, 11, 90])
sp.freeze_panes = "A8"

# ------------------------------------------------------------------ DIA A DIA
dd = wb.create_sheet("Dia a dia")
title(dd, "Presupuesto día a día · San Francisco · 28-sep → 4-nov-2026")
dd["A2"] = '="Escenario activo: "&Supuestos!B3&"   |   Cambia el escenario en Supuestos!B3. Cambia el alojamiento de cada noche con el desplegable de la columna D."'
dd["A2"].font = Font(italic=True, color="555555")
header(dd, 4, ["Fecha", "Día", "Semana", "Dónde duermes esta noche", "Alojamiento $", "Desayuno", "Almuerzo", "Cena", "Transporte", "Otros", "Total del día", "Tope política del día", "Margen vs tope", "Notas"])
dvh = DataValidation(type="list", formula1='"Airbnb,Casa Joaquín,Sin noche"', allow_blank=False); dd.add_data_validation(dvh)
for i, d in enumerate(DAYS):
    r = FIRST + i
    week = (d - START).days // 7 + 1
    dow = d.weekday()
    dd.cell(r, 1, d).number_format = DATEF
    dd.cell(r, 2, DIAS_ES[dow])
    dd.cell(r, 3, f"S{week}")
    if d == END: hous = "Sin noche"
    elif d >= JOAQUIN_FROM: hous = "Casa Joaquín"
    else: hous = "Airbnb"
    dd.cell(r, 4, hous).fill = in_fill; dvh.add(f"D{r}")
    dd.cell(r, 5, f'=IF(D{r}="Airbnb",Supuestos!$E$9*(1+Supuestos!$E$10),0)')
    dd.cell(r, 6, "=Supuestos!$E$11"); dd.cell(r, 7, "=Supuestos!$E$12"); dd.cell(r, 8, "=Supuestos!$E$13")
    if d == START: dd.cell(r, 9, "=Supuestos!$E$14+Supuestos!$E$19")
    elif d == END: dd.cell(r, 9, "=Supuestos!$E$14+Supuestos!$E$20")
    else: dd.cell(r, 9, "=Supuestos!$E$14")
    dd.cell(r, 10, "=Supuestos!$E$15+Supuestos!$E$21" if d == START else "=Supuestos!$E$15")
    dd.cell(r, 11, f"=SUM(E{r}:J{r})")
    dd.cell(r, 12, f'=IF(D{r}="Airbnb",Supuestos!$B$9,0)+Supuestos!$B$11+Supuestos!$B$12+Supuestos!$B$13+Supuestos!$B$14')
    dd.cell(r, 13, f"=L{r}-K{r}")
    notes = []
    if d == START: notes.append("Llegada SFO. Check-in Airbnb 3PM (SoMa). Traslado + eSIM.")
    if d == JOAQUIN_FROM: notes.append("SUPUESTO: check-out Airbnb y te mudas a casa de Joaquín (10 noches). Si son otras fechas, cambia la columna D.")
    if d == date(2026, 10, 5): notes.append("Deadline RAMP: reconciliar cargos de sep (28–30).")
    if d == date(2026, 11, 1): notes.append("Termina el horario de verano en EE.UU. (PST).")
    if d == END: notes.append("Vuelo de regreso. Sin noche. Tiquete flex: si vuelves antes, pon 'Sin noche' en los días que sobren y déjalos en 0.")
    if dow >= 5 and not notes: notes.append("Fin de semana")
    dd.cell(r, 14, " ".join(notes))
    for col in range(5, 14): dd.cell(r, col).number_format = USD
    for col in range(1, 15): dd.cell(r, col).border = box
    if dow >= 5:
        for col in range(1, 5): dd.cell(r, col).fill = sub_fill if col != 4 else in_fill
    dd.cell(r, 11).font = Font(bold=True)
tr = LAST + 1
dd.cell(tr, 1, "TOTAL VIAJE (sin vuelo; lavandería y comidas con clientes solo en Resumen)").font = Font(bold=True)
for col in range(5, 14):
    L = get_column_letter(col)
    c = dd.cell(tr, col, f"=SUM({L}{FIRST}:{L}{LAST})"); c.number_format = USD; c.font = Font(bold=True); c.fill = tot_fill; c.border = box
dd.cell(tr + 1, 1, "Promedio por día").font = Font(bold=True)
for col in range(5, 14):
    L = get_column_letter(col)
    c = dd.cell(tr + 1, col, f"=AVERAGE({L}{FIRST}:{L}{LAST})"); c.number_format = USD; c.fill = tot_fill; c.border = box
dd.conditional_formatting.add(f"M{FIRST}:M{LAST}", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=RED)))
dd.conditional_formatting.add(f"M{FIRST}:M{LAST}", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=PatternFill("solid", fgColor=GREEN)))
widths(dd, [13, 11, 8, 20, 14, 11, 11, 11, 12, 10, 14, 16, 14, 70])
dd.freeze_panes = "E5"

# ------------------------------------------------------------------ RESUMEN
rs = wb.create_sheet("Resumen", 0)
title(rs, "Presupuesto viaje San Francisco · German Tatis · 28-sep → 4-nov-2026", "Comparación: tope de la Travel & Expense Policy de Yuno (v1.0, abr-2026) vs costo real estimado en SF. USD.")
params = [
    ("Fecha de salida (BOG → SFO)", START, DATEF),
    ("Fecha de regreso (SFO → BOG)", END, DATEF),
    ("Días en viaje", f"=COUNT({DIA}!A{FIRST}:A{LAST})", "0"),
    ("Noches totales", "=B6-1", "0"),
    ("Noches en casa de Joaquín (sin costo)", f'=COUNTIF({DIA}!D{FIRST}:D{LAST},"Casa Joaquín")', "0"),
    ("Noches en Airbnb (pagadas)", f'=COUNTIF({DIA}!D{FIRST}:D{LAST},"Airbnb")', "0"),
    ("Escenario activo en 'Dia a dia'", "=Supuestos!B3", None),
    ("Tu parte del Airbnb", "=Supuestos!B5", PCT),
]
for i, (k, v, f) in enumerate(params):
    r = 4 + i
    rs.cell(r, 1, k).font = Font(bold=True); c = rs.cell(r, 2, v)
    if f: c.number_format = f
    c.fill = tot_fill; c.border = box
rs["D4"] = "Airbnb elegido"; rs["D4"].font = Font(bold=True)
rs["E4"] = "SOMA Executive Retreat Lower · 2 habitaciones / 1 baño · Clementina St (8th & 9th), SoMa · compartido con Samuel · check-in 28-sep 3PM"
rs["D5"] = "Aprobación"; rs["D5"].font = Font(bold=True)
rs["E5"] = "Justo + Sean OK en DM (17-sep, '<$4k/mes'). Pendiente: hilo en #travel-requests para el Airbnb y aviso a Finance por estadía 7+ noches (Policy §4)."
rs["D6"] = "Precio real del listing"; rs["D6"].font = Font(bold=True)
rs["E6"] = "NO CONFIRMADO. Airbnb no muestra el precio sin sesión. Actualiza Supuestos!C8 con el precio por noche real y todo se recalcula."
for r in (4, 5, 6): rs.cell(r, 5).alignment = Alignment(wrap_text=True, vertical="top")

header(rs, 13, ["Categoría", "Política (tope máximo)", "Realista (mercado SF)", "Ahorro", "Escenario activo (Día a día)", "Margen Realista vs Política", "Cómo se calcula"])
S = "Supuestos!"
def trip(col):  # col letter in Supuestos for a scenario
    return {
        "Vuelo (ya pagado)": f"={S}{col}18",
        "Alojamiento Airbnb (tu parte, sin impuestos)": f"=$B$9*{S}{col}9",
        "Impuestos Airbnb (TOT + TID si < 30 noches)": f"=$B$9*{S}{col}9*{S}{col}10",
        "Desayunos": f"=$B$6*{S}{col}11",
        "Almuerzos": f"=$B$6*{S}{col}12",
        "Cenas": f"=$B$6*{S}{col}13",
        "Transporte local (diario)": f"=$B$6*{S}{col}14",
        "Traslados aeropuerto (llegada + salida)": f"={S}{col}19+{S}{col}20",
        "eSIM / datos móviles": f"={S}{col}21",
        "Otros diarios": f"=$B$6*{S}{col}15",
        "Lavandería / varios únicos": f"={S}{col}22",
        "Comidas con clientes / equipo (opcional)": f"={S}B26",
    }
cats = list(trip("B").keys())
how = {
    "Vuelo (ya pagado)": "Fijo. Avianca Flex, comprado.",
    "Alojamiento Airbnb (tu parte, sin impuestos)": "Noches Airbnb × tu parte por noche (Supuestos fila 9).",
    "Impuestos Airbnb (TOT + TID si < 30 noches)": "16,25% sobre el alojamiento si reservas < 30 noches; 0% con 30+.",
    "Desayunos": "Días × valor por escenario.", "Almuerzos": "Días × valor por escenario.", "Cenas": "Días × valor por escenario.",
    "Transporte local (diario)": "Días × valor por escenario.",
    "Traslados aeropuerto (llegada + salida)": "En Política = 0 porque caben en el tope de $55/día.",
    "eSIM / datos móviles": "Reembolsable con recibo (Policy §9).",
    "Otros diarios": "Bolsillo propio salvo propósito de negocio.",
    "Lavandería / varios únicos": "Una sola vez.",
    "Comidas con clientes / equipo (opcional)": "Nº comidas × $85 (Supuestos B24). Fuera de los topes individuales.",
}
for i, cat in enumerate(cats):
    r = 14 + i
    rs.cell(r, 1, cat)
    rs.cell(r, 2, trip("B")[cat]); rs.cell(r, 3, trip("C")[cat]); rs.cell(r, 4, trip("D")[cat])
    rs.cell(r, 6, f"=B{r}-C{r}"); rs.cell(r, 7, how[cat])
    for col in (2, 3, 4, 5, 6): rs.cell(r, col).number_format = USD; rs.cell(r, col).border = box
    rs.cell(r, 1).border = box; rs.cell(r, 7).alignment = Alignment(wrap_text=True, vertical="top")
# active column: from daily sheet
rs["E14"] = f"={S}E18"
rs["E15"] = f"=SUMPRODUCT(({DIA}!D{FIRST}:D{LAST}=\"Airbnb\")*1)*{S}E9"
rs["E16"] = f"=SUM({DIA}!E{FIRST}:E{LAST})-E15"
rs["E17"] = f"=SUM({DIA}!F{FIRST}:F{LAST})"; rs["E18"] = f"=SUM({DIA}!G{FIRST}:G{LAST})"; rs["E19"] = f"=SUM({DIA}!H{FIRST}:H{LAST})"
rs["E20"] = f"=SUM({DIA}!I{FIRST}:I{LAST})-{S}E19-{S}E20"
rs["E21"] = f"={S}E19+{S}E20"; rs["E22"] = f"={S}E21"
rs["E23"] = f"=SUM({DIA}!J{FIRST}:J{LAST})-{S}E21"
rs["E24"] = f"={S}E22"; rs["E25"] = f"={S}B26"
TR = 14 + len(cats)  # 26
rs.cell(TR, 1, "TOTAL DEL VIAJE").font = Font(bold=True)
for col in (2, 3, 4, 5, 6):
    L = get_column_letter(col); c = rs.cell(TR, col, f"=SUM({L}14:{L}{TR-1})")
    c.number_format = USD; c.font = Font(bold=True); c.fill = tot_fill; c.border = box
rs.cell(TR + 1, 1, "Total sin vuelo").font = Font(bold=True)
rs.cell(TR + 2, 1, "Promedio por día (sin vuelo)").font = Font(bold=True)
rs.cell(TR + 3, 1, "Promedio por semana (sin vuelo)").font = Font(bold=True)
rs.cell(TR + 4, 1, "Total por mes de RAMP: septiembre (28–30 sep)").font = Font(bold=True)
rs.cell(TR + 5, 1, "Total por mes de RAMP: octubre").font = Font(bold=True)
rs.cell(TR + 6, 1, "Total por mes de RAMP: noviembre (1–4 nov)").font = Font(bold=True)
for col in (2, 3, 4, 5):
    L = get_column_letter(col)
    rs.cell(TR + 1, col, f"={L}{TR}-{L}14").number_format = USD
    rs.cell(TR + 2, col, f"={L}{TR+1}/$B$6").number_format = USD
    rs.cell(TR + 3, col, f"={L}{TR+2}*7").number_format = USD
    for cc in range(1, 4): rs.cell(TR + cc, col).border = box; rs.cell(TR + cc, col).fill = tot_fill
# RAMP months only for active scenario (from daily sheet)
rs.cell(TR + 4, 5, f'=SUMPRODUCT((MONTH({DIA}!A{FIRST}:A{LAST})=9)*{DIA}!K{FIRST}:K{LAST})').number_format = USD
rs.cell(TR + 5, 5, f'=SUMPRODUCT((MONTH({DIA}!A{FIRST}:A{LAST})=10)*{DIA}!K{FIRST}:K{LAST})').number_format = USD
rs.cell(TR + 6, 5, f'=SUMPRODUCT((MONTH({DIA}!A{FIRST}:A{LAST})=11)*{DIA}!K{FIRST}:K{LAST})').number_format = USD
rs.cell(TR + 4, 7, "Reconciliar en RAMP antes del 5-oct."); rs.cell(TR + 5, 7, "Reconciliar antes del 5-nov."); rs.cell(TR + 6, 7, "Reconciliar antes del 5-dic.")
for cc in (4, 5, 6): rs.cell(TR + cc, 5).fill = tot_fill; rs.cell(TR + cc, 5).border = box

# weekly table
WR = TR + 9
rs.cell(WR - 1, 1, "Por semana (escenario activo)").font = Font(bold=True, size=12, color=NAVY)
header(rs, WR, ["Semana", "Desde", "Hasta", "Días", "Total semana", "Tope política semana", "Margen"])
for w in range(1, 7):
    r = WR + w
    ws_ = START + timedelta((w - 1) * 7); we_ = min(ws_ + timedelta(6), END)
    rs.cell(r, 1, f"S{w}"); rs.cell(r, 2, ws_).number_format = DATEF; rs.cell(r, 3, we_).number_format = DATEF
    rs.cell(r, 4, f'=COUNTIF({DIA}!C{FIRST}:C{LAST},A{r})')
    rs.cell(r, 5, f'=SUMIF({DIA}!C{FIRST}:C{LAST},A{r},{DIA}!K{FIRST}:K{LAST})').number_format = USD
    rs.cell(r, 6, f'=SUMIF({DIA}!C{FIRST}:C{LAST},A{r},{DIA}!L{FIRST}:L{LAST})').number_format = USD
    rs.cell(r, 7, f"=F{r}-E{r}").number_format = USD
    for col in range(1, 8): rs.cell(r, col).border = box

# key takeaways
KR = WR + 9
rs.cell(KR, 1, "Lecturas clave").font = Font(bold=True, size=12, color=NAVY)
notes = [
    "1. Alojamiento: con el Airbnb compartido (~$130/noche la unidad, tu parte ~$65) estás muy por debajo del tope de $220/noche por persona. Es el mayor ahorro del viaje.",
    "2. Impuestos: si la reserva de Airbnb es de 27 noches paga 16,25% de impuestos de SF (TOT 14% + TID 2,25%). 30 noches sin impuesto cuestan MENOS que 27 con impuesto (27 × 1,1625 = 31,4 noches equivalentes). Evalúa reservar 30 noches.",
    "3. Comida: los topes de Yuno ($15 / $25 / $50 = $90/día, tip incluido) son realistas para SF. Un almuerzo casual sale $20–33 con tax + SF Mandate + tip; una cena mid-range $35–55. El desayuno es donde más se ahorra (mercado en el Airbnb).",
    "4. Transporte: el tope de $55/día es generoso. Con Muni pass 'A' ($104/mes, incluye BART dentro de SF) + 3–4 Uber por semana gastas ~$15/día. BART cubre SFO ($11,15) y la política lo pide como primera opción.",
    "5. Estadía 7+ noches: la política exige contactar a Finance antes de reservar. La aprobación en el DM (Justo/Sean) no sustituye el hilo en #travel-requests: cualquier cargo de hotel/Airbnb sin hilo aprobado se descuenta de nómina (Policy §2, §7).",
    "6. RAMP: reconciliar cada cargo (recibo + memo + categoría) antes del 5 del mes siguiente: 5-oct (sep), 5-nov (oct), 5-dic (nov). Comidas con clientes: etiquetar a todos los asistentes.",
    "7. Tiquete flex: si vuelves antes, marca 'Sin noche' en los días que sobren y pon las comidas en 0, o borra las filas. Los totales se recalculan solos.",
]
for i, n in enumerate(notes):
    c = rs.cell(KR + 1 + i, 1, n); c.alignment = Alignment(wrap_text=True, vertical="top")
    rs.merge_cells(start_row=KR + 1 + i, start_column=1, end_row=KR + 1 + i, end_column=7)
    rs.row_dimensions[KR + 1 + i].height = 32
widths(rs, [46, 20, 20, 14, 22, 22, 60])
rs.freeze_panes = "A4"

# ------------------------------------------------------------------ POLITICA
po = wb.create_sheet("Politica Yuno")
title(po, "Yuno Travel & Expense Policy v1.0 (abr-2026) · lo que aplica a este viaje", "Fuente: Google Doc 'Travel & Expense Policy' (Finance). Página de ciudad: San Francisco (Parte 3). Topes en USD, duros; exceder sin aprobación escrita = descuento de nómina.")
header(po, 4, ["Categoría (San Francisco)", "Tope", "Unidad", "Nota de la política"])
caps = [
    ("Hotel / Airbnb por noche (sin impuestos)", 220, "USD/noche", "3 estrellas equivalente. Airbnb permitido si cumple el tope. Estadías 7+ noches: contactar a Finance antes de reservar. Considerar Oakland/South Bay si es mucho más barato y BART es práctico."),
    ("Desayuno", 15, "USD/comida", "Recibo autogenerado si < $30. Tip incluido en el tope."),
    ("Almuerzo", 25, "USD/comida", "Recibo autogenerado si < $30. Tip incluido en el tope."),
    ("Cena", 50, "USD/comida", "Recibo obligatorio. Una bebida con la comida es aceptable dentro del tope. Tip incluido."),
    ("Comida con equipo / cliente, por persona", 85, "USD/persona", "Etiquetar a todos los asistentes en RAMP + propósito de negocio + recibo siempre."),
    ("Transporte local por día", 55, "USD/día", "Uber por defecto; transporte público primero donde exista. BART cubre SFO: siempre primera opción. Lyft solo si Uber no está disponible."),
    ("Tope total diario implícito (alojamiento + 3 comidas + transporte)", "=B5+B6+B7+B8+B10", "USD/día", "Suma de los topes individuales: $220 + $15 + $25 + $50 + $55 = $365. Referencia GSA FY2026 para SF: $272 lodging + $92 M&IE = $364."),
]
for i, (a, b, u, n) in enumerate(caps):
    r = 5 + i
    po.cell(r, 1, a); po.cell(r, 2, b).number_format = USD; po.cell(r, 3, u); po.cell(r, 4, n).alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(1, 5): po.cell(r, col).border = box
header(po, 13, ["Regla global", "Qué significa para este viaje", "Estado / acción"])
rules = [
    ("Aprobación en #travel-requests antes de reservar (§2)", "Vuelo y alojamiento deben tener hilo aprobado por Lex o Joha. El DM con Justo/Sean no es el hilo oficial.", "Vuelo: hilo del 14-sep OK. Airbnb: abrir hilo con el link del listing antes de pagar."),
    ("Reserva con 10+ días de anticipación (§2, §3)", "Salida 28-sep. La reserva del Airbnb debería estar hecha antes del 18-sep; si no, pedir aprobación escrita del Region Leader.", "Verificar."),
    ("Estadías de 7+ noches (§4)", "Contactar a Finance antes de reservar. Aplica de lleno (27 noches en Airbnb).", "Escribir a Finance / #travel-requests."),
    ("Todo con tarjeta RAMP (§1, §7)", "Nada con tarjeta personal salvo autorización escrita de Finance. Cargos personales: avisar y reembolsar en 5 días hábiles.", "Usar RAMP para Airbnb, comidas, transporte, eSIM."),
    ("Topes por comida, no por día (§5)", "$15 desayuno / $25 almuerzo / $50 cena. No se pueden sumar (p. ej. saltar desayuno y gastar $40 en almuerzo no aplica).", "Presupuesto construido por comida."),
    ("Tips dentro del tope (§5, NA)", "En EE.UU. el tip (18–20%) va dentro del tope, no encima. Un almuerzo de $25 en menú excede el tope tras tax + tip.", "Apuntar a ~$19 en menú para almuerzo y ~$38 en cena."),
    ("Reconciliación antes del 5 del mes siguiente (§8)", "Recibo + memo con propósito + categoría (+ asistentes en comidas grupales).", "5-oct, 5-nov, 5-dic."),
    ("Transporte: público primero, Uber por defecto (§6)", "BART desde SFO; Muni/BART dentro de SF; Uber cuando no sea práctico. Car rental requiere Region Leader.", "Muni pass 'A' $104/mes + BART SFO."),
    ("Viaje internacional (§9)", "Seguro de viaje lo da la empresa (verificar con People). SIM/datos reembolsables con recibo. Sin adelantos de efectivo.", "Confirmar seguro con People antes del 28-sep."),
    ("Bleisure (§10)", "Días personales corren por tu cuenta. Si extiendes, avisar a Region Leader y Finance. Con tiquete flex, volver antes no tiene costo de política.", "N/A por ahora."),
    ("No reembolsable", "Minibar, room service, pay-per-view, alcohol como ítem dominante, upgrades de seguro en alquiler de autos, car service privado.", "—"),
]
for i, (a, b, c) in enumerate(rules):
    r = 14 + i
    po.cell(r, 1, a); po.cell(r, 2, b); po.cell(r, 3, c)
    for col in range(1, 4): po.cell(r, col).border = box; po.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")
widths(po, [48, 60, 48, 80])
po.freeze_panes = "A5"

# ------------------------------------------------------------------ RESEARCH
re_ = wb.create_sheet("Research SF")
title(re_, "Research de costos reales en San Francisco (sep-2026)", "Fuentes públicas consultadas el 22-sep-2026. Los rangos alimentan las columnas Realista y Ahorro de Supuestos.")
header(re_, 4, ["Ítem", "Valor / rango", "Cómo lo usé", "Fuente"])
research = [
    ("Airbnb SF: tarifa promedio por noche", "$276 promedio; mediana $186; cuartil inferior $121; top 25% $314", "Ahorro = $121 (unidad). Referencia para validar el listing.", "https://www.airroi.com/airbnb-data/united-states/california/san-francisco"),
    ("Airbnb SF: ingreso mensual mediano por listing", "$3,301 / mes", "Consistente con el '<$4k/mes' de Sean para el listing elegido.", "https://www.airroi.com/airbnb-data/united-states/california/san-francisco"),
    ("Apartamentos amoblados por mes (Airbnb-friendly buildings)", "$2,030 – $4,244 / mes (NEMA $3,168; HQ $2,310; 77 Bluxome $2,030; Azure $4,244)", "Alternativa si el Airbnb cae. Con 30+ noches no hay TOT.", "https://www.airbnb.com/airbnb-friendly/apartments/san-francisco-ca"),
    ("Hotel 3 estrellas SF por noche", "$160 – $189 promedio; fines de semana hasta $332", "Comparación vs tope $220. Un hotel 3* solo cabe entre semana.", "https://www.booking.com/threestars/city/us/san-francisco.html"),
    ("Impuestos alojamiento < 30 noches", "TOT 14% + Tourism Improvement District 2,25% (Zona 1)", "Supuestos fila 10. 30+ noches = exento.", "https://sftreasurer.org/business/taxes-fees/transient-occupancy-tax-tot"),
    ("Airbnb: fee de servicio al huésped 2026", "Migró al host (15,5%); el huésped ve precio total sin fee separado", "No añadí fee de servicio aparte.", "https://www.airbnb.com/help/article/1857"),
    ("GSA per diem SF FY2026 (gobierno EE.UU.)", "Lodging $272 / noche; M&IE $92 / día", "Benchmark externo para validar topes Yuno ($220 / $90).", "https://www.gsa.gov/travel/plan-book/per-diem-rates/per-diem-rates-results?action=perdiems_report&city=san+Francisco&fiscal_year=2026&state=CA"),
    ("Café / desayuno", "Café $4–6 (latte $5,55); pastry/bagel $5–7", "Realista desayuno $12.", "https://www.expatistan.com/cost-of-living/san-francisco"),
    ("Almuerzo casual", "$15–25 en menú; combo fast-food ~$12", "Realista almuerzo $25 con tax + tip.", "https://www.worldcostofliving.com/city/san-francisco-usa"),
    ("Cena mid-range", "Plato $20–40; comida para dos $59–120; 3 tiempos ~$100", "Realista cena $45.", "https://www.expatistan.com/cost-of-living/san-francisco"),
    ("Recargos en restaurantes", "Sales tax 8,625% + SF Mandate (health) 3–7% + tip 17–22% = ~30% sobre el menú", "Explica por qué el tope de $25 se alcanza con un plato de ~$19.", "https://www.sfchronicle.com/food/restaurants/article/surcharge-fee-san-francisco-19561318.php"),
    ("Mercado / groceries (1 persona)", "$400–500 / mes; Trader Joe's es la más barata (SF Chronicle mar-2026)", "Ahorro desayuno $5 y cena $20.", "https://www.sfchronicle.com/projects/2026/cheapest-sf-grocery-prices-tariffs/"),
    ("Muni (bus, metro, cable car)", "$2,85 por viaje (120 min); pass mensual M $86; pass A (Muni + BART dentro de SF) $104", "Realista transporte $15/día (pass A + Ubers). Ahorro $6/día.", "https://www.sfmta.com/getting-around/muni/fares"),
    ("BART SFO ↔ downtown", "$11,15 una vía", "Traslado de salida (y llegada en Ahorro).", "https://www.bart.gov/tickets"),
    ("UberX SFO → downtown", "$35–55 normal; $65–90 con surge; +$5,50 airport surcharge", "Traslado de llegada Realista $45 (con maletas, check-in 3PM).", "https://getridewise.com/airports/sfo"),
    ("Uber intra-ciudad", "$20–50 por trayecto en hora pico", "3–4 viajes/semana × ~$25 en Realista.", "https://www.citizendailypost.com/faq/how-much-does-it-cost-to-get-around-in-san-francisco"),
    ("eSIM / datos", "T-Mobile US Pass $50/30 días (ilimitado, activar en EE.UU.); $25/7 días. Airalo 50 GB/30 días $42", "Realista $75 (30 + 7 días). Ahorro $57.", "https://www.airalo.com/united-states-esim"),
    ("Lavandería wash & fold", "$2,39 – $2,50 / lb; mínimo $15–25", "Varios $20 (el listing tiene lavadora).", "https://thelaundryhubsf.com/laundry-pricing-san-francisco"),
    ("Coworking day pass (si lo necesitas)", "$14 – $75 / día (Mindspace $42)", "No incluido: no es categoría de la política; pedir aprobación previa.", "https://wezoo.com/insights/coworking-day-pass"),
    ("Costo de vida general SF", "Comida 30,5% sobre promedio EE.UU. ($522/mes individual)", "Contexto.", "https://www.salary.com/research/cost-of-living/san-francisco-ca"),
]
for i, (a, b, c, d) in enumerate(research):
    r = 5 + i
    re_.cell(r, 1, a); re_.cell(r, 2, b); re_.cell(r, 3, c); re_.cell(r, 4, d)
    for col in range(1, 5): re_.cell(r, col).border = box; re_.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")
widths(re_, [40, 55, 50, 70])
re_.freeze_panes = "A5"

# ------------------------------------------------------------------ COMPARACION
cp = wb.create_sheet("Comparacion")
title(cp, "Política Yuno vs GSA vs mercado real de San Francisco (por día / por noche)", "Positivo en 'Margen' = el tope de Yuno alcanza. Negativo = el mercado supera el tope.")
header(cp, 4, ["Categoría", "Tope Yuno SF", "GSA FY2026 (benchmark)", "Mercado SF: rango bajo", "Mercado SF: rango alto", "Realista usado", "Margen tope vs realista", "Comentario"])
comp = [
    ("Alojamiento por noche (tu parte, sin impuestos)", 220, 272, 60.5, 186, "=Supuestos!C9", "Compartiendo 2BR quedas a ~30% del tope. Solo (entire place mediana $186) también cabe; hotel 3* fin de semana ($332) no."),
    ("Desayuno", 15, 20, 8, 15, "=Supuestos!C11", "Cabe. Café + pastry con tip ronda $12. GSA reparte $92 M&IE aprox. 20/25/40 + incidentales."),
    ("Almuerzo", 25, 25, 20, 33, "=Supuestos!C12", "Justo. Con tax + mandate + tip un plato de $25 en menú sale ~$33. Apuntar a ≤$19 en menú."),
    ("Cena", 50, 42, 35, 55, "=Supuestos!C13", "Cabe en mid-range. Una cena con entrada + plato + bebida en SoMa/Mission supera $50."),
    ("Transporte por día", 55, 5, 6, 30, "=Supuestos!C14", "Muy holgado. GSA da $5 de incidentales; aquí Muni pass + Uber ocasional ~$15."),
    ("Total por día (alojamiento + comidas + transporte)", "=SUM(B5:B9)", "=SUM(C5:C9)", "=SUM(D5:D9)", "=SUM(E5:E9)", "=SUM(F5:F9)", "Tope Yuno $365/día vs GSA $364/día: alineados. Tu gasto realista ronda $160/día con el Airbnb compartido."),
]
for i, (a, b, c, d, e, f, g) in enumerate(comp):
    r = 5 + i
    cp.cell(r, 1, a); cp.cell(r, 2, b); cp.cell(r, 3, c); cp.cell(r, 4, d); cp.cell(r, 5, e); cp.cell(r, 6, f)
    cp.cell(r, 7, f"=B{r}-F{r}"); cp.cell(r, 8, g)
    for col in range(2, 8): cp.cell(r, col).number_format = USD
    for col in range(1, 9): cp.cell(r, col).border = box; cp.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")
    if i == len(comp) - 1:
        for col in range(1, 9): cp.cell(r, col).font = Font(bold=True); cp.cell(r, col).fill = tot_fill
cp.conditional_formatting.add("G5:G10", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=RED)))
cp.conditional_formatting.add("G5:G10", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=PatternFill("solid", fgColor=GREEN)))
cp["A13"] = "Totales del viaje por escenario (38 días, 27 noches Airbnb + 10 en casa de Joaquín)"; cp["A13"].font = Font(bold=True, size=12, color=NAVY)
header(cp, 14, ["Escenario", "Total viaje (con vuelo)", "Total sin vuelo", "Promedio por día sin vuelo", "vs tope de política"])
for i, (name, col) in enumerate([("Política (tope máximo)", "B"), ("Realista (mercado SF)", "C"), ("Ahorro", "D")]):
    r = 15 + i
    cp.cell(r, 1, name); cp.cell(r, 2, f"=Resumen!{col}26"); cp.cell(r, 3, f"=Resumen!{col}27"); cp.cell(r, 4, f"=Resumen!{col}28"); cp.cell(r, 5, f"=Resumen!B26-B{r}")
    for c_ in range(2, 6): cp.cell(r, c_).number_format = USD
    for c_ in range(1, 6): cp.cell(r, c_).border = box
widths(cp, [46, 16, 20, 20, 20, 16, 20, 80])

wb.save(OUT)
print("saved", OUT)
