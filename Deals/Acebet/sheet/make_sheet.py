# -*- coding: utf-8 -*-
"""Acebet proposal explainer workbook (27-oct-2026): 4 tabs, live formulas, house format. Upload as xlsx -> Google Sheet."""
import sys; sys.path.insert(0, '../build')
from model import CUTS, RATES, PLATFORM, var, total
TX = 680000; V, T = var(TX), total(TX)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
NAVY, F = '10284A', 'Roboto'
thin = Side(style='thin', color='000000'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
H = dict(font=Font(name=F, size=10, bold=True, color='FFFFFF'), fill=PatternFill('solid', fgColor=NAVY),
         alignment=Alignment(horizontal='left', vertical='center', wrap_text=True))
INPUT = PatternFill('solid', fgColor='FFF2CC'); KEY = PatternFill('solid', fgColor='E8EAF5')
wb = Workbook()
def sheet(ws, title, widths):
    ws.title = title; ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
def row(ws, r, vals, header=False, bold=False, fmt=None, fill=None):
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v); cell.border = BOX
        if header:
            for k, s in H.items(): setattr(cell, k, s)
        else:
            cell.font = Font(name=F, size=10, bold=bold or c == 1); cell.alignment = Alignment(vertical='top', wrap_text=True)
            if fmt and c > 1 and fmt[c - 2]: cell.number_format = fmt[c - 2]
            if fill: cell.fill = fill
    if header: ws.row_dimensions[r].height = 30

# ---------------- 1. Summary
ws = wb.active; sheet(ws, 'Summary', [30, 70, 60])
row(ws, 1, ['Topic', 'What we are proposing', 'Why / how to say it'], header=True)
S = [
 ('Client', 'Acebet: US sweepstakes / social casino operator, Cyprus entity. Contact: Boris (co-founder).', 'Referred by Coinflow. First call 27-Oct-2026.'),
 ('What they want', 'Split volume across several acquirers by fixed percentages, cascade declines, add acquirers fast.', 'Leverage on price and redundancy. Approval on Breeze is already in the 90s, so this is not an uplift sale.'),
 ('Volume basis', 'About 680,000 successful transactions a month, average ticket about $33.', 'Real last 30 days: 678,498 deposits, $22.77M, all on Breeze (internal data; never cite it to the client). Boris said ~700K and ~$25M.'),
 ('Pricing structure', f'Platform fee ${PLATFORM:,}/month flat + a rate per successful transaction in 4 tranches.', 'Tranche (ladder) pricing: each tranche billed at its own rate within the month, so the average falls as volume grows.'),
 ('Tranches', ' / '.join(f'{"0" if i == 0 else f"{CUTS[i-1]+1:,}"} to {CUTS[i]:,}: ${RATES[i]:.3f}' if CUTS[i] else f'above {CUTS[i-1]:,}: ${RATES[i]:.3f}' for i in range(4)), 'Rates per German (27-Oct): 3, 2.5, 2 and 1.5 cents, platform fee on top. Top tranche sits at Spreedly\'s quoted ~1.5 cents.'),
 ('Today\'s bill', f'See tab Calculator: ${V:,.0f} usage + ${PLATFORM:,} platform = ${T:,.0f} a month (${T*12:,.0f} a year).', f'Average ${V/TX:.4f} per transaction before the platform fee; ${T/TX:.4f} all-in.'),
 ('Billing rules', 'Charged only on approved transactions. A retried or cascaded payment counts once. Declines cost nothing. No setup fee.', 'Matches what was promised on the call (retries charged once when accepted).'),
 ('Included', 'Orchestration and rules engine; volume split, cascading, retries; monitors and failover; hosted fields, SDKs, sandbox; new provider connections maintained by Yuno; dashboard and insights; KAM, TAM, solutions engineer.', '"Every acquirer you add is a routing change, not a new fee."'),
 ('Not included (pending review)', 'Reconciliation across Breeze, Coinflow, Whop and Checkout, for further review.', 'Boris showed interest in reconciliation; scope and price later. Payouts removed from the proposal.'),
 ('Terms', '36-month contract. Rates locked, tranches reviewed annually. Pricing valid 30 days, subject to final agreement.', 'Minimums and billing start date not stated yet.'),
 ('Competitive frame', 'PaymentIQ ~$0.04/trx, Spreedly ~$0.015/trx (as quoted by Boris). See tab Vs competitors.', f'Our usage average (${V/TX:.3f}) sits between them; all-in (${T/TX:.3f}) is about PaymentIQ. Unknown whether their quotes include platform fees: ask.'),
 ('Open risks (internal)', 'Pricing approval (Justo + Salesforce). Wallet cascading across MoRs to confirm with Ben Lee. 2-week acquirer SLA not in writing yet. Whop needs a provider ticket. Legal check on Cyprus entity.', 'Policy floor is $0.01 per transaction; every tranche is above it (lowest $0.015).'),
]
for i, r in enumerate(S, 2): row(ws, i, r)
ws.freeze_panes = 'B2'

# ---------------- 2. Calculator (live)
ws = wb.create_sheet(); sheet(ws, 'Calculator', [34, 18, 18, 16, 18, 18])
row(ws, 1, ['Input', 'Value', '', '', '', ''], header=True)
row(ws, 2, ['Successful transactions / month (edit)', 680000, '', '', '', ''], fmt=['#,##0', None, None, None, None], fill=None)
ws['B2'].fill = INPUT
row(ws, 3, ['Platform fee / month', PLATFORM, '', '', '', ''], fmt=['$#,##0', None, None, None, None]); ws['B3'].fill = INPUT
row(ws, 5, ['Tranche', 'From', 'To', 'Rate / trx', 'Trx in tranche', 'Cost / month'], header=True)
for i in range(4):
    r = 6 + i; lo = 0 if i == 0 else CUTS[i - 1]; hi = CUTS[i]
    row(ws, r, [f'Tranche {i+1}', lo + (1 if i else 0), hi if hi else 'no cap', RATES[i],
                f'=MAX(0,MIN($B$2,{hi})-{lo})' if hi else f'=MAX(0,$B$2-{lo})', f'=E{r}*D{r}'],
        fmt=['#,##0', '#,##0', '$0.000', '#,##0', '$#,##0'])
    ws[f'D{r}'].fill = INPUT
row(ws, 11, ['Usage subtotal', '', '', '', '=SUM(E6:E9)', '=SUM(F6:F9)'], bold=True, fmt=[None, None, None, '#,##0', '$#,##0'])
row(ws, 12, ['Platform fee', '', '', '', '', '=B3'], fmt=[None, None, None, None, '$#,##0'])
row(ws, 13, ['Total / month', '', '', '', '', '=F11+F12'], bold=True, fmt=[None, None, None, None, '$#,##0'], fill=KEY)
row(ws, 14, ['Total / year', '', '', '', '', '=F13*12'], fmt=[None, None, None, None, '$#,##0'])
row(ws, 15, ['Average usage rate / trx (before platform fee)', '', '', '', '', '=F11/B2'], bold=True, fmt=[None, None, None, None, '$0.0000'], fill=KEY)
row(ws, 16, ['All-in rate / trx (with platform fee)', '', '', '', '', '=F13/B2'], bold=True, fmt=[None, None, None, None, '$0.0000'], fill=KEY)
ws['A18'] = 'Yellow cells are inputs: change the volume or a rate and every number updates.'; ws['A18'].font = Font(name=F, size=9, italic=True)

# ---------------- 3. Growth
ws = wb.create_sheet(); sheet(ws, 'Growth', [26, 16, 16, 16, 16, 16])
row(ws, 1, ['Monthly transactions', 'Usage / month', 'Platform fee', 'Total / month', 'Usage avg / trx', 'All-in / trx'], header=True)
def usage(c):
    parts = []
    for i in range(4):
        lo = 0 if i == 0 else CUTS[i - 1]; hi = CUTS[i]
        parts.append(f'MAX(0,MIN({c},{hi})-{lo})*{RATES[i]}' if hi else f'MAX(0,{c}-{lo})*{RATES[i]}')
    return '=' + '+'.join(parts)
for j, tx in enumerate([680000, 1000000, 1500000, 2000000, 3000000], 2):
    row(ws, j, [tx, usage(f'A{j}'), f'=Calculator!$B$3', f'=B{j}+C{j}', f'=B{j}/A{j}', f'=D{j}/A{j}'],
        fmt=['$#,##0', '$#,##0', '$#,##0', '$0.0000', '$0.0000'])
    ws[f'A{j}'].number_format = '#,##0'
ws['A8'] = 'Volumes above today are illustrative (expansion to Canada, Australia, France, Spain), not forecasts.'; ws['A8'].font = Font(name=F, size=9, italic=True)

# ---------------- 4. Vs competitors
ws = wb.create_sheet(); sheet(ws, 'Vs competitors', [30, 16, 18, 18, 50])
row(ws, 1, ['Option', 'Rate / trx', 'Cost / month', 'Cost / year', 'Note'], header=True)
row(ws, 2, ['Spreedly (quoted)', 0.015, '=B2*Calculator!$B$2', '=C2*12', 'As quoted by Boris; unknown if a platform fee applies on top.'], fmt=['$0.000', '$#,##0', '$#,##0', None])
row(ws, 3, ['Yuno usage only', '=Calculator!F15', '=Calculator!F11', '=C3*12', 'Tranche usage before the platform fee.'], fmt=['$0.0000', '$#,##0', '$#,##0', None])
row(ws, 4, ['Yuno all-in', '=Calculator!F16', '=Calculator!F13', '=C4*12', 'Usage plus the $10,000 platform fee.'], fmt=['$0.0000', '$#,##0', '$#,##0', None], fill=KEY)
row(ws, 5, ['PaymentIQ (quoted)', 0.04, '=B5*Calculator!$B$2', '=C5*12', 'As quoted by Boris; unknown if a platform fee applies on top.'], fmt=['$0.000', '$#,##0', '$#,##0', None])
ws['A7'] = 'Ask Boris whether the PaymentIQ and Spreedly quotes are all-in; the comparison changes if they carry platform fees.'; ws['A7'].font = Font(name=F, size=9, italic=True)
wb.save('Acebet_proposal_explainer.xlsx'); print('saved')
