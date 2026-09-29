#!/usr/bin/env python3
"""Cruz Verde: solicitud de información como Google Sheet, en formato de formulario.

Cada hoja tiene a la izquierda el campo que pedimos y al lado el espacio (amarillo) para
la respuesta de Cruz Verde. El contenido sale de build_data_request.py (la misma fuente
del Excel): aquí solo se cambia la presentación, sin columnas calculadas, totales ni
filas de ejemplo.

Requisito: el Sheet debe estar compartido como Editor con la service account
gtm-claude-editor@gtm-claude-tools-260922.iam.gserviceaccount.com (~/.config/gsuite/sa.json).

Uso: python3 build_data_request_gsheet.py            (falla si el Sheet ya tiene hojas creadas)
     python3 build_data_request_gsheet.py --rebuild  (borra las hojas y las vuelve a crear)
"""
import math
import os
import sys
import tempfile
import warnings

warnings.filterwarnings("ignore")

from google.oauth2 import service_account  # noqa: E402
from googleapiclient.discovery import build as gbuild  # noqa: E402

sys.argv, ARGS = sys.argv[:1], sys.argv[1:]
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_data_request as src  # noqa: E402

SPREADSHEET_ID = "1xipYBM6dNEVmCj_cxezIAqsfMb-gbPCimMJmXaGIMJQ"
SA_KEY = os.path.expanduser("~/.config/gsuite/sa.json")

BLUE, LILAC, BLACK, DGRAY = src.BLUE, src.LILAC, src.BLACK, src.DGRAY
INPUT, WHITE, LINE = src.INPUT, src.WHITE, src.LINE

NCOLS = 8                                   # A..H
WIDTHS = [330, 125, 125, 125, 125, 125, 125, 300]
NOTE_OVERRIDES = {
    "3.3": "Los segmentos no son excluyentes entre sí.",
    "5.2": "Si ya tienen su propia separación entre rechazos definitivos y temporales, esta es la tabla para escribirla.",
    "8.1": "Las columnas de días y de costo son promedios por medio de pago.",
}
HEADER_OVERRIDES = {("5.2", "Según Cruz Verde: rechazos (#)"): "Rechazos (#)"}
CONTENIDO = {
    "1. Export": "Alternativa más rápida: los 24 campos del export transaccional, una fila por intento de pago.",
    "2. Volumen": "Volumen mensual, por canal, por rango de ticket y por tipo de entrega. Proyección 2027.",
    "3. Mix de medios": "Volumen y aprobación por medio de pago, marca y tipo de tarjeta. Efectivo y links de pago.",
    "4. Aprobación": "Aprobación por banco emisor, por intento frente a por orden, reintentos y 3DS.",
    "5. Rechazos": "Motivos de rechazo y separación entre rechazos definitivos (hard) y temporales (soft).",
    "6. PSE y efectivo": "Embudo de PSE, PSE por banco y embudo de pagos en efectivo.",
    "7. Incidentes": "Caídas o intermitencias y disponibilidad del servicio.",
    "8. Devoluciones": "Devoluciones por medio de pago, motivos, proceso de PSE, contracargos y fraude.",
    "9. Costos": "Tarifas por medio de pago, otros costos y contrato actual.",
    "10. Conciliación": "Proceso actual de conciliación.",
    "11. Checkout": "Embudo de checkout, plataforma e integración.",
    "12. Nuevos medios": "Medios de pago de interés, evidencia de demanda y planes.",
}


# ---------- 1. contenido: se graba lo que build_data_request.py le pide a cada hoja ----------
TABS = []


class Rec(src.Sheet):
    def __init__(self, wb, name, num, title, subtitle, grid=None):
        super().__init__(wb, name, num, title, subtitle, grid)
        self.rec = {"name": name, "num": num, "title": title, "subtitle": subtitle, "blocks": []}
        TABS.append(self.rec)

    def table(self, title, cols, rows, prio, area, uso, exp, dato=None, note=None, example=None, total=False):
        out = super().table(title, cols, rows, prio, area, uso, exp, dato=dato, note=note,
                            example=example, total=total)
        norm = [(c[0], c[1], c[2] if len(c) > 2 else 1, c[3] if len(c) > 3 else None) for c in cols]
        self.rec["blocks"].append(dict(kind="table", id=out[0], title=title, cols=norm, rows=rows, prio=prio,
                                       area=area, exp=exp, dato=dato, note=note))
        return out

    def questions(self, title, items, note=None):
        k0 = self.k
        super().questions(title, items, note)
        self.rec["blocks"].append(dict(kind="questions", title=title, items=[
            dict(id=f"{self.num}.{k0 + i + 1}", q=q, prio=prio, dato=dato)
            for i, (q, prio, area, uso, exp, dato) in enumerate(items)]))


def load_content():
    src.Sheet = Rec
    src.OUT = os.path.join(tempfile.mkdtemp(), "fuente.xlsx")
    path, n = src.build()
    from openpyxl import load_workbook
    ws = load_workbook(path)["1. Export"]
    campos = [[c.value for c in row[:6]] for row in ws.iter_rows() if isinstance(row[0].value, int)]
    return campos, n


# ---------- 2. celdas ----------
def rgb(h):
    return {"red": int(h[0:2], 16) / 255, "green": int(h[2:4], 16) / 255, "blue": int(h[4:6], 16) / 255}


def cell(value=None, bg=WHITE, bold=False, italic=False, color=BLACK, size=10, h="LEFT", v="TOP",
         border=True, fmt=None, options=None, link=None):
    tf = {"fontFamily": "Arial", "fontSize": size, "bold": bold, "italic": italic, "foregroundColor": rgb(color)}
    if link:
        tf["link"] = {"uri": link}
    f = {"backgroundColor": rgb(bg), "textFormat": tf, "horizontalAlignment": h, "verticalAlignment": v,
         "wrapStrategy": "WRAP"}
    if border:
        b = {"style": "SOLID", "color": rgb(LINE)}
        f["borders"] = {"top": b, "bottom": b, "left": b, "right": b}
    if fmt == "int":
        f["numberFormat"] = {"type": "NUMBER", "pattern": "#,##0"}
    elif fmt == "pct":
        f["numberFormat"] = {"type": "PERCENT", "pattern": "0.0%"}
    d = {"userEnteredFormat": f}
    if value is not None:
        d["userEnteredValue"] = ({"numberValue": value} if isinstance(value, (int, float))
                                 else {"stringValue": str(value)})
    if options:
        d["dataValidation"] = {"condition": {"type": "ONE_OF_LIST",
                                             "values": [{"userEnteredValue": o} for o in options]},
                               "showCustomUi": True, "strict": False}
    return d


def lines(text, px, size=10):
    if not text:
        return 1
    cpl = max(1, int((px - 8) / (size * 0.62)))
    return sum(max(1, math.ceil(len(p) / cpl)) for p in str(text).split("\n"))


class Tab:
    """Grilla en memoria de una hoja; al final se convierte en requests de la API."""

    def __init__(self, gid, name, widths, title, subtitle, index):
        self.gid, self.name, self.widths, self.index = gid, name, widths, index
        self.n = len(widths)
        self.rows, self.heights, self.merges = [], [], []
        self.band(title, bg=BLUE, color=WHITE, size=14, height=42)
        self.band(subtitle, bg=LILAC, color=DGRAY, size=10, bold=False, height=28)
        self.blank()

    def px(self, c1, c2):
        return sum(self.widths[c1:c2 + 1])

    def add(self, cells, height):
        """cells: lista de (col_inicio, col_fin, celda)."""
        r = len(self.rows)
        row = [cell(border=False) for _ in range(self.n)]
        for c1, c2, d in cells:
            row[c1] = d
            for c in range(c1 + 1, c2 + 1):
                row[c] = {"userEnteredFormat": d["userEnteredFormat"]}
            if c2 > c1:
                self.merges.append((r, c1, c2))
        self.rows.append(row)
        self.heights.append(int(height))
        return r

    def band(self, text, bg=LILAC, color=BLACK, size=11, bold=True, height=30):
        self.add([(0, self.n - 1, cell(text, bg=bg, bold=bold, color=color, size=size, v="MIDDLE",
                                       border=False))], height)

    def note(self, text):
        self.add([(0, self.n - 1, cell(text, italic=True, color=DGRAY, size=9, v="MIDDLE", border=False))],
                 max(22, lines(text, self.px(0, self.n - 1), 9) * 15 + 8))

    def blank(self, height=14):
        self.add([], height)

    def requests(self):
        g = self.gid
        req = [{"addSheet": {"properties": {
            "sheetId": g, "title": self.name, "index": self.index, "tabColor": rgb(BLUE),
            "gridProperties": {"rowCount": len(self.rows) + 3, "columnCount": self.n, "frozenRowCount": 2,
                               "hideGridlines": True}}}}]
        req.append({"updateCells": {"rows": [{"values": r} for r in self.rows],
                                    "fields": "userEnteredValue,userEnteredFormat,dataValidation",
                                    "start": {"sheetId": g, "rowIndex": 0, "columnIndex": 0}}})
        for r, c1, c2 in self.merges:
            req.append({"mergeCells": {"mergeType": "MERGE_ALL", "range": {
                "sheetId": g, "startRowIndex": r, "endRowIndex": r + 1,
                "startColumnIndex": c1, "endColumnIndex": c2 + 1}}})
        for i, w in enumerate(self.widths):
            req.append({"updateDimensionProperties": {
                "range": {"sheetId": g, "dimension": "COLUMNS", "startIndex": i, "endIndex": i + 1},
                "properties": {"pixelSize": w}, "fields": "pixelSize"}})
        for i, h in enumerate(self.heights):
            req.append({"updateDimensionProperties": {
                "range": {"sheetId": g, "dimension": "ROWS", "startIndex": i, "endIndex": i + 1},
                "properties": {"pixelSize": h}, "fields": "pixelSize"}})
        return req


HEAD = dict(bg=BLUE, bold=True, color=WHITE, h="CENTER", v="MIDDLE")


def render_table(t, b):
    info = f"Prioridad {b['prio']}. Área sugerida: {b['area']}."
    if b["exp"] == "Sí":
        info += " Si nos envían el export transaccional, esta tabla la calculamos nosotros."
    elif b["exp"] == "Parcial":
        info += " Con el export transaccional la calculamos en parte."
    t.band(f"{b['id']}  {b['title']}")
    t.note(info)
    if b["dato"]:
        t.note(f"Lo que nos han compartido, por confirmar: {b['dato']}.")
    note = NOTE_OVERRIDES.get(b["id"], b["note"])
    if note:
        t.note(note)

    # columnas: el campo va en A; las respuestas van de B en adelante; el último texto llega hasta H
    keep = [(j, c) for j, c in enumerate(b["cols"]) if c[1] != "calc"]
    pos, c0 = {}, 1
    for n, (j, (head, kind, span, opt)) in enumerate(keep):
        if n == 0:
            pos[j] = (0, 0)
        elif n == len(keep) - 1 and kind == "text":
            pos[j] = (c0, NCOLS - 1)
        else:
            w = span if kind in ("text", "fixed") else 1
            pos[j] = (c0, c0 + w - 1)
            c0 += w
    assert c0 <= NCOLS, b["title"]

    hl = 1
    cells = []
    for j, (head, kind, span, opt) in keep:
        head = HEADER_OVERRIDES.get((b["id"], head), head)
        a, z = pos[j]
        cells.append((a, z, cell(head, **HEAD)))
        hl = max(hl, lines(head, t.px(a, z)))
    t.add(cells, max(32, hl * 16 + 10))

    for row in b["rows"]:
        if isinstance(row, tuple) and row[0] == "CALC":
            continue
        cells, hl = [], 1
        for j, (head, kind, span, opt) in keep:
            a, z = pos[j]
            v = row[j] if j < len(row) else None
            given = v is not None
            if kind == "label":
                d = cell(v, bg=WHITE if given else INPUT)
            elif kind == "fixed":
                d = cell(v, color=DGRAY, size=9)
            elif kind == "int":
                d = cell(v, bg=INPUT, h="RIGHT", fmt="int")
            elif kind == "pct":
                d = cell(v, bg=INPUT, h="RIGHT", fmt="pct")
            elif kind == "yn":
                d = cell(v, bg=INPUT, options=["Sí", "No"])
            elif kind == "pick":
                d = cell(v, bg=WHITE if given else INPUT, options=opt["options"])
            else:  # text: si viene diligenciado es una referencia, no una respuesta
                d = cell(v, bg=WHITE, color=DGRAY, size=9) if given else cell(None, bg=INPUT)
            cells.append((a, z, d))
            if isinstance(v, str):
                hl = max(hl, lines(v, t.px(a, z), 9 if kind in ("fixed", "text") else 10))
        t.add(cells, max(26, hl * 16 + 8))
    t.blank()


def render_questions(t, b):
    t.band(b["title"])
    t.add([(0, 0, cell("Campo", **HEAD)), (1, 4, cell("Respuesta de Cruz Verde", **HEAD)),
           (5, 5, cell("Prioridad", **HEAD)),
           (6, 7, cell("Lo que nos han compartido, por confirmar", **HEAD))], 32)
    for it in b["items"]:
        q = f"{it['id']}  {it['q']}"
        hl = max(lines(q, t.px(0, 0)), lines(it["dato"], t.px(6, 7), 9), 2)
        t.add([(0, 0, cell(q)), (1, 4, cell(None, bg=INPUT)),
               (5, 5, cell(it["prio"], bold=True, h="CENTER")),
               (6, 7, cell(it["dato"], italic=True, color=DGRAY, size=9))], hl * 16 + 10)
    t.blank()


def tab_export(gid, index, campos):
    w = [40, 240, 380, 290, 210, 80, 110, 300]
    t = Tab(gid, "1. Export", w, "1. Export transaccional (alternativa más rápida)",
            f"Una fila por intento de pago. Período: {src.PERIODO}. Formato: CSV o Excel. "
            "En esta hoja solo hay que indicar qué campos pueden compartir.", index)
    for n in [
        "Si nos comparten este export, nosotros calculamos buena parte de las tablas de las demás hojas.",
        "Incluir todos los intentos de pago del canal digital: aprobados, rechazados, pendientes, cancelados y "
        "devueltos. Los rechazados son los más importantes para el análisis de aprobación.",
        "La columna \"Campo de referencia en Mercado Pago\" usa los nombres de la API de pagos de Mercado Pago "
        "como guía. El nombre puede variar según el reporte o la integración que tengan; lo importante es el dato.",
        "Sin datos personales: no necesitamos nombres, documentos, correos ni números completos de tarjeta.",
    ]:
        t.note(n)
    t.blank()
    heads = ["#", "Campo", "Descripción", "Campo de referencia en Mercado Pago", "Ejemplo de formato",
             "Prioridad", "¿Disponible?", "Comentarios de Cruz Verde"]
    t.add([(j, j, cell(h, **HEAD)) for j, h in enumerate(heads)], 36)
    for num, campo, desc, ref, ej, prio in campos:
        hl = max(lines(desc, w[2]), lines(campo, w[1]), lines(ref, w[3], 9))
        t.add([(0, 0, cell(num, color=DGRAY, h="CENTER")), (1, 1, cell(campo)), (2, 2, cell(desc)),
               (3, 3, cell(ref, color=DGRAY, size=9)), (4, 4, cell(str(ej), italic=True, color=DGRAY, size=9)),
               (5, 5, cell(prio, bold=True, h="CENTER")),
               (6, 6, cell(None, bg=INPUT, h="CENTER", options=["Sí", "No", "Parcial"])),
               (7, 7, cell(None, bg=INPUT))], max(26, hl * 16 + 8))
    return t


def tab_inicio(gid, gids):
    w = [230, 780]
    t = Tab(gid, "Inicio", w, "Cruz Verde + Yuno",
            "Información para el business case. Versión 1, 29 de septiembre de 2026.", 0)

    def kv(k, v, kbg=WHITE, link=None):
        t.add([(0, 0, cell(k, bold=True, bg=kbg, color=src.BLUE if link else BLACK, link=link)), (1, 1, cell(v))],
              max(26, max(lines(v, w[1]), lines(k, w[0])) * 16 + 8))

    for k, v in [
        ("Objetivo",
         "Construir el business case de Cruz Verde con las cifras reales de su operación de pagos y no con "
         "supuestos de mercado. Cada dato alimenta una parte específica del caso: aprobación de pagos, nuevos "
         "medios de pago, costo transaccional, conciliación, devoluciones y continuidad del servicio."),
        ("Cómo llenar",
         "Cada hoja tiene a la izquierda el campo que necesitamos y al lado el espacio para su respuesta. "
         "Solo hay que completar las celdas en amarillo."),
        ("Una alternativa más rápida",
         "Si nos comparten el export transaccional descrito en la hoja \"1. Export\", nosotros calculamos las "
         "tablas que así lo indican en su descripción y ustedes completan solo el resto."),
        ("Período",
         f"Últimos 12 meses cerrados: {src.PERIODO}. Si no es posible, un mínimo de 6 meses. "
         "Cuando una tabla pide un período distinto, lo indica en su descripción."),
        ("Moneda y formato",
         "Pesos colombianos (COP), sin decimales. Cantidades: 1.000. Montos: 45.000.000. Porcentajes: 90%."),
        ("Alcance",
         "Canal digital de Cruz Verde Colombia: sitio web, app y links de pago. Las preguntas sobre tiendas "
         "físicas lo indican de forma explícita."),
        ("Prioridad",
         "A: imprescindible para el business case.\nB: deseable, mejora la precisión del caso.\n"
         "C: si está disponible."),
        ("Lo que nos han compartido",
         "En varios campos incluimos la cifra que nos dieron en la llamada del 21 de mayo o en la reunión del "
         "25 de septiembre. La idea es que la confirmen o la corrijan con el dato real."),
        ("Si un dato no existe",
         "Escribir \"No disponible\". En esos casos acordamos con ustedes un supuesto y lo dejamos declarado en "
         "el business case. Las cifras aproximadas también sirven: basta con indicar \"aprox.\"."),
        ("Confidencialidad",
         "La información se comparte una vez esté firmado el acuerdo de confidencialidad, que está en trámite. "
         "No necesitamos datos personales de clientes ni números completos de tarjeta: los primeros 6 dígitos "
         "(BIN) son suficientes para identificar banco y marca."),
        ("Área sugerida",
         "Cada tabla indica el área que probablemente tiene el dato (Ecommerce, Finanzas, TI). Es solo una guía."),
        ("Contacto en Yuno", "German Tatis, german.tatis@y.uno"),
    ]:
        kv(k, v)
    t.blank()
    t.band("Contenido")
    for name, desc in CONTENIDO.items():
        kv(name, desc, link=f"#gid={gids[name]}")
    t.blank()
    t.band("Definiciones")
    for k, v in [
        ("Intento de pago", "Cada vez que un cliente envía un pago para autorización. Una misma orden puede tener varios intentos."),
        ("Orden", "Pedido del cliente en el sistema de Cruz Verde, sin importar cuántos intentos de pago tuvo."),
        ("Tasa de aprobación por intento", "Pagos aprobados dividido entre intentos de pago."),
        ("Tasa de aprobación por orden", "Órdenes pagadas dividido entre órdenes con al menos un intento de pago."),
        ("Rechazo definitivo (hard decline)",
         "El banco emisor indica que esa tarjeta no puede aprobar el pago: tarjeta vencida, cancelada, reportada "
         "o inactiva. Reintentar con la misma tarjeta no cambia el resultado."),
        ("Rechazo temporal (soft decline)",
         "El rechazo depende del momento o de la ruta: fondos insuficientes, error técnico, rechazo genérico. "
         "Puede aprobarse con un reintento, por otro procesador o más tarde."),
        ("Rechazo por riesgo", "El pago lo detiene el motor de prevención de fraude del proveedor o del banco."),
        ("Rechazo por datos del comprador", "El cliente digitó mal un dato de la tarjeta. Se corrige en el checkout."),
        ("Devolución", "Reintegro total o parcial de un pago aprobado, iniciado por Cruz Verde."),
        ("Contracargo", "Disputa de un pago iniciada por el tarjetahabiente ante su banco."),
        ("Ventas aprobadas", "Monto total en COP de los pagos aprobados en el período."),
    ]:
        kv(k, v)
    return t


# ---------- 3. publicación ----------
def main():
    campos, n = load_content()
    assert len(campos) == 24 and TABS[0]["name"] == "1. Export"
    gids = {name: 1000 + i for i, name in enumerate(CONTENIDO, 1)}
    tabs = [tab_inicio(1000, gids), tab_export(gids["1. Export"], 1, campos)]
    for i, rec in enumerate(TABS[1:], 2):
        t = Tab(gids[rec["name"]], rec["name"], WIDTHS, rec["title"],
                rec["subtitle"] + " Completar las celdas en amarillo.", i)
        for b in rec["blocks"]:
            (render_table if b["kind"] == "table" else render_questions)(t, b)
        tabs.append(t)

    creds = service_account.Credentials.from_service_account_file(
        SA_KEY, scopes=["https://www.googleapis.com/auth/spreadsheets"])
    api = gbuild("sheets", "v4", credentials=creds, cache_discovery=False).spreadsheets()
    meta = api.get(spreadsheetId=SPREADSHEET_ID, fields="sheets.properties(sheetId,title)").execute()
    old = [s["properties"] for s in meta["sheets"]]
    built = [p for p in old if p["sheetId"] >= 1000]
    if built and "--rebuild" not in ARGS:
        raise SystemExit("El Sheet ya tiene las hojas creadas. Usar --rebuild para borrarlas y recrearlas "
                         "(se pierde lo que Cruz Verde haya escrito).")

    TMP = 999
    req = [{"updateSpreadsheetProperties": {"properties": {"locale": "es_CO", "timeZone": "America/Bogota"},
                                            "fields": "locale,timeZone"}},
           {"addSheet": {"properties": {"sheetId": TMP, "title": "_tmp"}}}]
    req += [{"deleteSheet": {"sheetId": p["sheetId"]}} for p in old]
    api.batchUpdate(spreadsheetId=SPREADSHEET_ID, body={"requests": req}).execute()
    for t in tabs:
        api.batchUpdate(spreadsheetId=SPREADSHEET_ID, body={"requests": t.requests()}).execute()
        print(f"  {t.name}: {len(t.rows)} filas")
    api.batchUpdate(spreadsheetId=SPREADSHEET_ID, body={"requests": [{"deleteSheet": {"sheetId": TMP}}]}).execute()
    print(f"OK https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/edit ({n} solicitudes, {len(tabs)} hojas)")


if __name__ == "__main__":
    main()
