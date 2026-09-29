#!/usr/bin/env python3
"""Cruz Verde: solicitud de información para el business case (xlsx).

Genera "Solicitud de información - Cruz Verde + Yuno.xlsx" en Deals/Cruz Verde/.
Todo el contenido vive en este script: cada tabla y cada pregunta se registra
también en la hoja "Resumen", así que las dos vistas nunca se desalinean.

Uso: python3 build_data_request.py [ruta_salida.xlsx]
"""
import math
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    os.path.dirname(HERE), "Solicitud de información - Cruz Verde + Yuno.xlsx")

# Marca Yuno (Assets/Yuno_Brandbook)
BLUE, LILAC, BLACK, GRAY, DGRAY = "3E4FE0", "E8EAF5", "282A30", "92959B", "616366"
INPUT, CALC, EXFILL, WHITE, LINE = "FFF4CC", "F3F4FA", "EDEDED", "FFFFFF", "D5D8E6"

D_MAY = "llamada del 21 de mayo"
D_SEP = "reunión del 25 de septiembre"
PERIODO = "septiembre de 2025 a agosto de 2026"
MESES = ["sep-2025", "oct-2025", "nov-2025", "dic-2025", "ene-2026", "feb-2026",
         "mar-2026", "abr-2026", "may-2026", "jun-2026", "jul-2026", "ago-2026"]

F_INT = '#,##0;-#,##0;"-"'
F_PCT = "0.0%"

RESUMEN = []  # (id, hoja, texto, prioridad, área, uso, export, dato)


def font(size=10, bold=False, italic=False, color=BLACK):
    return Font(name="Arial", size=size, bold=bold, italic=italic, color=color)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


_s = Side(style="thin", color=LINE)
BORDER = Border(left=_s, right=_s, top=_s, bottom=_s)


def n_lines(text, width, size=10):
    if text is None or text == "":
        return 1
    cpl = max(1.0, width * (10.0 / size) * 1.25)
    return sum(max(1, math.ceil(len(p) / cpl)) for p in str(text).split("\n"))


def row_height(lines, size=10):
    return max(18.0, lines * size * 1.38 + 6)


# ---------- fórmulas ----------
class Cx:
    def __init__(self, letters, r, r1, r2, i):
        self.letters, self.r, self.r1, self.r2, self.i = letters, r, r1, r2, i

    def c(self, idx):
        return self.letters[idx]


def rate(num, den):
    """num / den de la misma fila; vacío si falta algún dato."""
    def f(cx):
        n, d = f"{cx.c(num)}{cx.r}", f"{cx.c(den)}{cx.r}"
        return f'=IF(AND(ISNUMBER({n}),ISNUMBER({d})),IF({d}>0,{n}/{d},""),"")'
    return f


def share(idx):
    """Participación de la fila sobre el total de la columna."""
    def f(cx):
        col = cx.c(idx)
        rng = f"{col}${cx.r1}:{col}${cx.r2}"
        return f'=IF(ISNUMBER({col}{cx.r}),IF(SUM({rng})>0,{col}{cx.r}/SUM({rng}),""),"")'
    return f


def over_first(idx):
    """Valor de la fila sobre el de la primera fila de la tabla (embudo)."""
    def f(cx):
        a, b = f"{cx.c(idx)}{cx.r}", f"{cx.c(idx)}${cx.r1}"
        return f'=IF(AND(ISNUMBER({a}),ISNUMBER({b})),IF({b}>0,{a}/{b},""),"")'
    return f


def over_prev(idx):
    """Valor de la fila sobre el de la fila anterior (embudo paso a paso)."""
    def f(cx):
        if cx.i == 0:
            return None
        a, b = f"{cx.c(idx)}{cx.r}", f"{cx.c(idx)}{cx.r - 1}"
        return f'=IF(AND(ISNUMBER({a}),ISNUMBER({b})),IF({b}>0,{a}/{b},""),"")'
    return f


# ---------- hoja ----------
class Sheet:
    GRID = [6, 44, 15, 15, 15, 15, 15, 15, 36]  # A..I

    def __init__(self, wb, name, num, title, subtitle, grid=None):
        self.ws = wb.create_sheet(name)
        self.name, self.num, self.k = name, num, 0
        self.grid = grid or self.GRID
        self.ncols = len(self.grid)
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = BLUE
        for i, w in enumerate(self.grid, 1):
            ws.column_dimensions[L(i)].width = w
        self.merge(1, 1, self.ncols, title, font(16, True, color=WHITE), fill(BLUE), h=34, border=False,
                   align=Alignment(vertical="center", indent=1))
        self.merge(2, 1, self.ncols, subtitle, font(10, color=DGRAY), fill(LILAC), h=24, border=False,
                   align=Alignment(vertical="center", indent=1, wrap_text=True))
        ws.freeze_panes = "A3"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.r = 4
        self.dvs = {}

    # -- utilidades
    def width(self, c1, c2):
        return sum(self.grid[c1 - 1:c2])

    def merge(self, r, c1, c2, text, fnt, fl=None, h=None, border=True, align=None, fmt=None):
        ws = self.ws
        if c2 > c1:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        cell = ws.cell(row=r, column=c1)
        cell.value = text
        cell.font = fnt
        cell.alignment = align or Alignment(wrap_text=True, vertical="top")
        if fmt:
            cell.number_format = fmt
        for c in range(c1, c2 + 1):
            x = ws.cell(row=r, column=c)
            if fl:
                x.fill = fl
            if border:
                x.border = BORDER
        if h:
            ws.row_dimensions[r].height = h
        return cell

    def grow(self, r, lines, size=10):
        cur = self.ws.row_dimensions[r].height or 0
        self.ws.row_dimensions[r].height = max(cur, row_height(lines, size))

    def validation(self, options, ref):
        key = ",".join(options)
        if key not in self.dvs:
            dv = DataValidation(type="list", formula1=f'"{key}"', allow_blank=True)
            dv.error = "Elegir una opción de la lista"
            self.ws.add_data_validation(dv)
            self.dvs[key] = dv
        self.dvs[key].add(ref)

    def next_id(self):
        self.k += 1
        return f"{self.num}.{self.k}"

    def band(self, text):
        self.merge(self.r, 1, self.ncols, text, font(11, True, color=BLACK), fill(LILAC), h=24, border=False,
                   align=Alignment(vertical="center", indent=1))
        self.r += 1

    def note(self, text, bold=False):
        w = self.width(1, self.ncols)
        self.merge(self.r, 1, self.ncols, text, font(9, bold, italic=not bold, color=DGRAY), border=False,
                   align=Alignment(wrap_text=True, vertical="top", indent=1))
        self.grow(self.r, n_lines(text, w - 4, 9), 9)
        self.r += 1

    def gap(self, n=1):
        self.r += n

    # -- tabla de datos
    def table(self, title, cols, rows, prio, area, uso, exp, dato=None, note=None,
              example=None, total=False):
        """cols: (encabezado, tipo, span, opciones). Tipos: label, int, pct, text, yn, pick, calc."""
        ws = self.ws
        tid = self.next_id()
        RESUMEN.append((tid, self.name, title, prio, area, uso, exp, dato or ""))
        self.band(f"{tid}  {title}")
        self.note(f"Prioridad {prio}. Área sugerida: {area}. Lo usamos para: {uso}.")
        if dato:
            self.note(f"Dato que tenemos hoy, a confirmar: {dato}.")
        if note:
            self.note(note)

        cols = [(c[0], c[1], c[2] if len(c) > 2 else 1, c[3] if len(c) > 3 else None) for c in cols]
        pos, c0 = [], 2
        for _, _, span, _ in cols:
            pos.append((c0, c0 + span - 1))
            c0 += span
        assert c0 - 1 <= self.ncols, f"{title}: la tabla no cabe en la grilla"
        letters = [L(a) for a, _ in pos]

        # encabezado
        hl = 1
        for (head, _, _, _), (a, b) in zip(cols, pos):
            self.merge(self.r, a, b, head, font(10, True, color=WHITE), fill(BLUE),
                       align=Alignment(wrap_text=True, vertical="center", horizontal="center"))
            hl = max(hl, n_lines(head, self.width(a, b) - 1))
        self.grow(self.r, hl)
        self.r += 1

        def put(r, j, value, style):
            head, kind, span, opt = cols[j]
            a, b = pos[j]
            num = kind in ("int", "pct") or (kind == "calc")
            fmt = None
            if kind == "int":
                fmt = F_INT
            elif kind == "pct":
                fmt = F_PCT
            elif kind == "calc":
                fmt = (opt or {}).get("fmt", F_PCT)
            if style == "example":
                fnt, fl = font(10, italic=True, color=DGRAY), fill(EXFILL)
            elif style == "total":
                fnt, fl = font(10, True), fill(CALC)
            elif kind == "calc":
                fnt, fl = font(10), fill(CALC)
            elif kind == "label" and value is not None:
                fnt, fl = font(10), fill(WHITE)
            elif kind == "fixed":
                fnt, fl = font(10, color=DGRAY), fill(WHITE)
            else:
                fnt, fl = font(10), fill(INPUT)
            al = Alignment(wrap_text=True, vertical="top",
                           horizontal="right" if num else "left")
            self.merge(r, a, b, value, fnt, fl, align=al, fmt=fmt)
            if isinstance(value, str) and not value.startswith("="):
                self.grow(r, n_lines(value, self.width(a, b) - 1))
            if style == "data" and kind == "yn":
                self.validation(["Sí", "No"], f"{L(a)}{r}")
            if style == "data" and kind == "pick":
                self.validation(opt["options"], f"{L(a)}{r}")

        if example:
            ws.cell(row=self.r, column=1, value="Ej.").font = font(9, italic=True, color=DGRAY)
            ws.cell(row=self.r, column=1).alignment = Alignment(horizontal="center", vertical="top")
            self.ws.row_dimensions[self.r].height = 18
            for j, v in enumerate(example):
                put(self.r, j, v, "example")
            self.r += 1

        r1, r2 = self.r, self.r + len(rows) - 1
        for i, row in enumerate(rows):
            self.ws.row_dimensions[self.r].height = 18
            cx = Cx(letters, self.r, r1, r2, i)
            if isinstance(row, tuple) and row[0] == "CALC":
                _, label, fn, fmt = row
                for j in range(len(cols)):
                    a, b = pos[j]
                    v = label if j == 0 else fn(cx) if j == 1 else None
                    self.merge(self.r, a, b, v, font(10, bold=(j == 1), italic=(j == 0)), fill(CALC),
                               align=Alignment(wrap_text=True, vertical="top",
                                               horizontal="right" if j == 1 else "left"),
                               fmt=fmt if j == 1 else None)
                self.r += 1
                continue
            for j, (head, kind, span, opt) in enumerate(cols):
                v = row[j] if j < len(row) else None
                if kind == "calc":
                    f = (opt or {}).get("rows", {}).get(i, (opt or {}).get("f"))
                    v = f(cx) if f else None
                put(self.r, j, v, "data")
            self.r += 1

        if total:
            self.ws.row_dimensions[self.r].height = 18
            cx = Cx(letters, self.r, r1, r2, -1)
            for j, (head, kind, span, opt) in enumerate(cols):
                if j == 0:
                    v = "Total"
                elif kind == "int" and not (opt or {}).get("nototal"):
                    v = f"=SUM({letters[j]}{r1}:{letters[j]}{r2})"
                elif kind == "calc" and (opt or {}).get("total"):
                    v = opt["f"](cx)
                else:
                    v = None
                put(self.r, j, v, "total")
            self.r += 1
        self.gap()
        return tid, r1, r2, letters

    # -- preguntas
    def questions(self, title, items, note=None):
        """items: (pregunta, prioridad, área, uso, export, dato)."""
        self.band(title)
        if note:
            self.note(note)
        heads = [(1, 1, "ID"), (2, 2, "Pregunta"), (3, 6, "Respuesta de Cruz Verde"),
                 (7, 7, "Prioridad"), (8, 9, "Dato que tenemos hoy, a confirmar")]
        for a, b, h in heads:
            self.merge(self.r, a, b, h, font(10, True, color=WHITE), fill(BLUE),
                       align=Alignment(wrap_text=True, vertical="center", horizontal="center"))
        self.ws.row_dimensions[self.r].height = 20
        self.r += 1
        for q, prio, area, uso, exp, dato in items:
            qid = self.next_id()
            RESUMEN.append((qid, self.name, q, prio, area, uso, exp, dato or ""))
            self.merge(self.r, 1, 1, qid, font(9, color=DGRAY), fill(WHITE),
                       align=Alignment(horizontal="center", vertical="top"))
            self.merge(self.r, 2, 2, q, font(10), fill(WHITE))
            self.merge(self.r, 3, 6, None, font(10), fill(INPUT))
            self.merge(self.r, 7, 7, prio, font(10, True), fill(WHITE),
                       align=Alignment(horizontal="center", vertical="top"))
            self.merge(self.r, 8, 9, dato, font(9, italic=True, color=DGRAY), fill(WHITE))
            lines = max(n_lines(q, self.width(2, 2) - 1), n_lines(dato, self.width(8, 9) - 1, 9), 2)
            self.grow(self.r, lines)
            self.r += 1
        self.gap()


# ---------- contenido ----------
def build():
    wb = Workbook()
    wb.remove(wb.active)

    # ===== Instrucciones =====
    ws = wb.create_sheet("Instrucciones")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = BLACK
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 120
    for r, (t, f_, fl_, h) in enumerate([
        ("Cruz Verde + Yuno", font(18, True, color=WHITE), fill(BLUE), 38),
        ("Solicitud de información para el business case. Versión 1, 29 de septiembre de 2026.",
         font(10, color=DGRAY), fill(LILAC), 24)], 1):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(row=r, column=1, value=t)
        c.font, c.alignment = f_, Alignment(vertical="center", indent=1)
        for col in (1, 2):
            ws.cell(row=r, column=col).fill = fl_
        ws.row_dimensions[r].height = h

    blocks = [
        ("Objetivo",
         "Construir el business case de Cruz Verde con las cifras reales de su operación de pagos y no con "
         "supuestos de mercado. Cada dato de este archivo alimenta una parte específica del caso: aprobación "
         "de pagos, nuevos medios de pago, costo transaccional, conciliación, devoluciones y continuidad del servicio."),
        ("Cómo compartir la información",
         "Hay dos caminos y se pueden combinar.\n"
         "Opción A (recomendada): export transaccional. Si nos comparten el detalle por intento de pago descrito en "
         "la hoja \"1. Export\", nosotros calculamos la mayoría de las tablas de las hojas 2 a 8. Ustedes solo "
         "responden lo que en la hoja \"Resumen\" aparece con \"No\" o \"Parcial\" en la columna \"Se calcula con el export\".\n"
         "Opción B: tablas resumen. Completar directamente las hojas 2 a 12 con cifras agregadas."),
        ("Período",
         f"Últimos 12 meses cerrados: {PERIODO}. Si no es posible, un mínimo de 6 meses. "
         "Cuando una tabla pide un período distinto, lo indica en su descripción."),
        ("Moneda", "Todas las cifras en pesos colombianos (COP), sin decimales."),
        ("Alcance",
         "Canal digital de Cruz Verde Colombia: sitio web, app y links de pago. Las preguntas sobre tiendas "
         "físicas lo indican de forma explícita."),
        ("Prioridad",
         "A: imprescindible para el business case.\nB: deseable, mejora la precisión del caso.\nC: si está disponible."),
        ("Dato que tenemos hoy",
         "En varias solicitudes incluimos la cifra que nos compartieron en la llamada del 21 de mayo o en la "
         "reunión del 25 de septiembre. La idea es que la confirmen o la corrijan con el dato real."),
        ("Si un dato no existe",
         "Escribir \"No disponible\". En esos casos acordamos con ustedes un supuesto y lo dejamos declarado en el "
         "business case. Las cifras aproximadas también sirven: basta con indicar \"aprox.\"."),
        ("Confidencialidad",
         "La información se comparte una vez esté firmado el acuerdo de confidencialidad, que está en trámite. "
         "No necesitamos datos personales de clientes ni números completos de tarjeta: los primeros 6 dígitos "
         "(BIN) son suficientes para identificar banco y marca."),
        ("Área sugerida",
         "Cada solicitud indica el área que probablemente tiene el dato (Ecommerce, Finanzas, TI). Es solo una guía."),
        ("Contacto en Yuno", "German Tatis, german.tatis@y.uno"),
    ]
    r = 4

    def kv(r, k, v, kfill=WHITE, vfont=None):
        a, b = ws.cell(row=r, column=1, value=k), ws.cell(row=r, column=2, value=v)
        a.font, b.font = font(10, True), vfont or font(10)
        a.fill = fill(kfill)
        a.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        a.border = b.border = BORDER
        ws.row_dimensions[r].height = row_height(max(n_lines(v, 119), n_lines(k, 28)))

    for k, v in blocks:
        kv(r, k, v)
        r += 1

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c = ws.cell(row=r, column=1, value="Cómo leer las celdas")
    c.font, c.alignment = font(11, True), Alignment(vertical="center", indent=1)
    for col in (1, 2):
        ws.cell(row=r, column=col).fill = fill(LILAC)
    ws.row_dimensions[r].height = 24
    r += 1
    for label, color, text, fnt in [
        ("Amarillo", INPUT, "Celdas para completar por Cruz Verde.", font(10)),
        ("Lila claro", CALC, "Celdas calculadas con fórmula. No hace falta modificarlas.", font(10)),
        ("Gris en cursiva", EXFILL,
         "Fila de ejemplo marcada con \"Ej.\". Solo muestra el formato esperado: las cifras son ilustrativas "
         "y no corresponden a Cruz Verde. Los totales no la incluyen.", font(10, italic=True, color=DGRAY)),
    ]:
        kv(r, label, text, kfill=color, vfont=fnt)
        r += 1
    kv(r, "Ejemplo de formato", "Cantidades: 1.000    Montos en COP: 45.000.000    Porcentajes: 90,0%", EXFILL,
       font(10, italic=True, color=DGRAY))
    r += 2

    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c = ws.cell(row=r, column=1, value="Definiciones")
    c.font, c.alignment = font(11, True), Alignment(vertical="center", indent=1)
    for col in (1, 2):
        ws.cell(row=r, column=col).fill = fill(LILAC)
    ws.row_dimensions[r].height = 24
    r += 1
    for k, v in [
        ("Intento de pago", "Cada vez que un cliente envía un pago para autorización. Una misma orden puede tener varios intentos."),
        ("Orden", "Pedido del cliente en el sistema de Cruz Verde, sin importar cuántos intentos de pago tuvo."),
        ("Tasa de aprobación por intento", "Pagos aprobados dividido entre intentos de pago."),
        ("Tasa de aprobación por orden", "Órdenes pagadas dividido entre órdenes con al menos un intento de pago."),
        ("Rechazo definitivo (hard decline)",
         "El banco emisor indica que esa tarjeta no puede aprobar el pago: tarjeta vencida, cancelada, reportada o "
         "inactiva. Reintentar con la misma tarjeta no cambia el resultado."),
        ("Rechazo temporal (soft decline)",
         "El rechazo depende del momento o de la ruta: fondos insuficientes, error técnico, rechazo genérico. "
         "Puede aprobarse con un reintento, por otro procesador o más tarde."),
        ("Rechazo por riesgo", "El pago lo detiene el motor de prevención de fraude del proveedor o del banco."),
        ("Rechazo por datos del comprador", "El cliente digitó mal un dato de la tarjeta. Se corrige en el checkout."),
        ("Devolución", "Reintegro total o parcial de un pago aprobado, iniciado por Cruz Verde."),
        ("Contracargo", "Disputa de un pago iniciada por el tarjetahabiente ante su banco."),
        ("Ventas aprobadas", "Monto total en COP de los pagos aprobados en el período."),
    ]:
        kv(r, k, v)
        r += 1
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # La hoja Resumen se crea aquí para que quede segunda; se llena al final.
    ws_res = wb.create_sheet("Resumen")

    # ===== 1. Export =====
    grid = [5, 36, 52, 38, 28, 11, 14, 36]
    s = Sheet(wb, "1. Export", 1, "1. Export transaccional (Opción A, recomendada)",
              f"Una fila por intento de pago. Período: {PERIODO}. Formato: CSV o Excel.", grid)
    RESUMEN.append(("1.1", "1. Export", "Export transaccional por intento de pago con los 24 campos de la hoja",
                    "A", "TI y Finanzas", "Todas las partes del caso", "Es el export", ""))
    for t in [
        "Incluir todos los intentos de pago del canal digital: aprobados, rechazados, pendientes, cancelados y devueltos. "
        "Los rechazados son los más importantes para el análisis de aprobación.",
        "La columna \"Campo de referencia en Mercado Pago\" usa los nombres de la API de pagos de Mercado Pago como guía. "
        "El nombre puede variar según el reporte o la integración que tengan; lo importante es el dato.",
        "Sin datos personales: no necesitamos nombres, documentos, correos ni números completos de tarjeta.",
    ]:
        s.note(t)
    s.gap()
    heads = ["#", "Campo", "Descripción", "Campo de referencia en Mercado Pago", "Ejemplo de formato",
             "Prioridad", "¿Disponible?", "Comentarios de Cruz Verde"]
    for j, h in enumerate(heads, 1):
        s.merge(s.r, j, j, h, font(10, True, color=WHITE), fill(BLUE),
                align=Alignment(wrap_text=True, vertical="center", horizontal="center"))
    s.ws.row_dimensions[s.r].height = 32
    s.r += 1
    CV = "Dato de Cruz Verde"
    campos = [
        ("Identificador del pago", "Identificador único de cada intento de pago.", "id", "1234567890", "A"),
        ("Identificador de la orden", "Número de la orden en el sistema de Cruz Verde. Permite agrupar los intentos de una misma compra.", "external_reference", "CV-000123", "A"),
        ("Fecha y hora de creación", "Momento en que el cliente envió el pago.", "date_created", "2026-08-15 14:32:05", "A"),
        ("Fecha y hora de aprobación", "Momento en que el pago quedó aprobado. Vacío si no se aprobó.", "date_approved", "2026-08-15 14:32:09", "B"),
        ("Monto de la transacción (COP)", "Valor cobrado al cliente.", "transaction_amount", "45000", "A"),
        ("Tipo de medio de pago", "Tarjeta de crédito, tarjeta de débito, PSE, efectivo u otro.", "payment_type_id", "credit_card", "A"),
        ("Medio o marca", "Marca de la tarjeta o medio específico.", "payment_method_id", "visa", "A"),
        ("Estado final", "Resultado del intento: aprobado, rechazado, pendiente, cancelado o devuelto.", "status", "rejected", "A"),
        ("Motivo o detalle del estado", "Motivo del rechazo o detalle del estado. Es el campo clave para separar rechazos definitivos y temporales.", "status_detail", "cc_rejected_insufficient_amount", "A"),
        ("Código de respuesta del banco emisor", "Código original que devolvió el banco, si el reporte lo trae.", "Dato del reporte, si existe", "51", "B"),
        ("BIN de la tarjeta", "Primeros 6 dígitos de la tarjeta. Identifica banco emisor, marca, tipo y país.", "card.first_six_digits", "450995", "A"),
        ("Banco emisor", "Banco que emitió la tarjeta.", "issuer_id", "Código o nombre del banco", "B"),
        ("Número de cuotas", "Cuotas elegidas por el cliente.", "installments", "1", "C"),
        ("Banco seleccionado en PSE", "Entidad financiera que el cliente eligió para pagar por PSE.", "transaction_details.financial_institution", "Código o nombre del banco", "B"),
        ("Identificador anónimo del cliente", "Código interno del cliente, sin datos personales. Permite medir reintentos y recurrencia.", "payer.id", "987654", "B"),
        ("Canal de venta", "Sitio web, app iOS, app Android, link de pago o televentas.", CV, "App Android", "A"),
        ("Tipo de entrega", "Mismo día, día siguiente o recogida en tienda.", CV, "Mismo día", "B"),
        ("Tarjeta guardada o digitada", "Si el cliente pagó con una tarjeta guardada o la digitó en esa compra.", CV, "Guardada", "C"),
        ("Monto devuelto (COP)", "Valor reintegrado al cliente, total o parcial.", "transaction_amount_refunded", "45000", "A"),
        ("Fecha de la devolución", "Fecha en que se ejecutó la devolución.", "Dato del reporte de devoluciones", "2026-08-18", "B"),
        ("Contracargo", "Si el pago tuvo una disputa del tarjetahabiente.", "Dato del reporte de contracargos", "No", "B"),
        ("Comisión cobrada (COP)", "Comisión que descontó el proveedor en esa transacción.", "fee_details", "1350", "A"),
        ("Monto neto recibido (COP)", "Valor que recibió Cruz Verde después de comisiones.", "transaction_details.net_received_amount", "43650", "A"),
        ("Fecha de liberación del dinero", "Fecha en que el dinero quedó disponible para Cruz Verde.", "money_release_date", "2026-08-16", "B"),
    ]
    for i, (campo, desc, ref, ej, prio) in enumerate(campos, 1):
        vals = [i, campo, desc, ref, ej, prio, None, None]
        for j, v in enumerate(vals, 1):
            inp = j in (7, 8)
            fnt = font(10, italic=(j == 5), color=DGRAY if j in (1, 4, 5) else BLACK, bold=(j == 6))
            al = Alignment(wrap_text=True, vertical="top", horizontal="center" if j in (1, 6, 7) else "left")
            s.merge(s.r, j, j, v, fnt, fill(INPUT if inp else WHITE), align=al)
        s.validation(["Sí", "No", "Parcial"], f"G{s.r}")
        s.grow(s.r, max(n_lines(desc, grid[2] - 1), n_lines(ref, grid[3] - 1), n_lines(campo, grid[1] - 1)))
        s.r += 1
    assert len(campos) == 24

    # ===== 2. Volumen =====
    s = Sheet(wb, "2. Volumen", 2, "2. Volumen y ventas del canal digital",
              f"Período: {PERIODO}. Cifras en COP.")
    s.table("Volumen mensual de todos los medios de pago",
            [("Mes", "label"), ("Intentos de pago (#)", "int"), ("Pagos aprobados (#)", "int"),
             ("Tasa de aprobación", "calc", 1, {"f": rate(2, 1), "total": True}),
             ("Ventas aprobadas (COP)", "int"),
             ("Ticket promedio (COP)", "calc", 1, {"f": rate(4, 2), "fmt": F_INT, "total": True}),
             ("Órdenes únicas (#)", "int"), ("Clientes únicos que pagaron (#)", "int", 1, {"nototal": True}),
             ],
            [[m] for m in MESES], "A", "Ecommerce y Finanzas", "la base del caso y el precio por tiers", "Sí",
            dato=f"más de 100.000 transacciones al mes, cerca de 80.000 por Mercado Pago y 20.000 en efectivo ({D_MAY})",
            example=["ene-2026", 1000, 900, 0.9, 45000000, 50000, 850, 700], total=True)
    s.table("Volumen por canal, acumulado de los 12 meses",
            [("Canal", "label"), ("Intentos de pago (#)", "int"), ("Pagos aprobados (#)", "int"),
             ("Tasa de aprobación", "calc", 1, {"f": rate(2, 1), "total": True}),
             ("Ventas aprobadas (COP)", "int"),
             ("Participación en ventas", "calc", 1, {"f": share(4)}),
             ("Comentarios", "text", 2)],
            [["Sitio web"], ["App iOS"], ["App Android"], ["Link de pago en punto de venta"],
             ["Televentas o call center"], ["Otro (especificar en comentarios)"]],
            "A", "Ecommerce", "aprobación por canal y alcance de la integración", "Parcial",
            example=["Sitio web", 1000, 900, 0.9, 45000000, 0.4, "Texto libre"], total=True)
    s.table("Distribución por rango de ticket, acumulado de los 12 meses",
            [("Rango de ticket (COP)", "label"), ("Pagos aprobados (#)", "int"),
             ("Participación en transacciones", "calc", 1, {"f": share(1)}),
             ("Ventas aprobadas (COP)", "int"),
             ("Intentos con tarjeta (#)", "int"), ("Aprobados con tarjeta (#)", "int"),
             ("Tasa de aprobación en tarjetas", "calc", 1, {"f": rate(5, 4), "total": True}),
             ("Comentarios", "text", 1)],
            [["Hasta 30.000"], ["30.001 a 60.000"], ["60.001 a 100.000"], ["100.001 a 200.000"],
             ["200.001 a 500.000"], ["Más de 500.000"]],
            "A", "Ecommerce y Finanzas", "el costo por transacción frente al ticket y la aprobación por monto", "Sí",
            dato=f"ticket promedio de COP 130.000 y a la baja, con 40% de las transacciones por debajo de COP 30.000 ({D_MAY})",
            example=["Hasta 30.000", 400, 0.4, 8000000, 180, 160, 0.889, "Texto libre"], total=True)
    s.table("Órdenes por tipo de entrega, acumulado de los 12 meses",
            [("Tipo de entrega", "label"), ("Órdenes pagadas (#)", "int"), ("Ventas aprobadas (COP)", "int"),
             ("Órdenes canceladas después del pago (#)", "int"),
             ("Medios de pago habilitados hoy", "text", 2), ("Medios de pago que quisieran habilitar", "text", 2)],
            [["Mismo día"], ["Día siguiente"], ["Recogida en tienda"], ["Otro (especificar)"]],
            "B", "Ecommerce", "el checkout con medios de pago distintos por tipo de entrega", "Parcial",
            dato=f"interés en un checkout con medios de pago distintos para entrega el mismo día y al día siguiente ({D_SEP})",
            example=["Mismo día", 500, 25000000, 20, "PSE y tarjetas", "PSE, tarjetas y billeteras"], total=False)
    s.questions("Preguntas sobre volumen y proyección", [
        ("Ticket promedio actual del canal digital (COP).", "A", "Ecommerce", "Base del caso", "Sí",
         f"COP 130.000 y a la baja ({D_MAY})."),
        ("Volumen mensual que esperan para 2027 (pagos aprobados por mes).", "A", "Ecommerce y Finanzas",
         "Precio por tiers", "No", f"Confirmaron que sí alcanzan las 100.000 transacciones mensuales ({D_SEP})."),
        ("Crecimiento esperado del canal digital en 2027 (%).", "B", "Ecommerce", "Proyección del caso", "No", None),
        ("Fechas de mayor demanda del año y número de transacciones del día pico.", "B", "Ecommerce",
         "Continuidad", "Sí", None),
        ("Transacciones por minuto en el momento de mayor demanda.", "C", "TI", "Continuidad", "Sí", None),
        ("Porcentaje de clientes que compraron más de una vez en los últimos 12 meses.", "B", "Ecommerce",
         "Nuevos medios de pago", "Parcial", None),
    ])

    # ===== 3. Mix de medios =====
    s = Sheet(wb, "3. Mix de medios", 3, "3. Mix de medios de pago",
              f"Período: {PERIODO}, acumulado. Cifras en COP.")
    mixcols = lambda first: [(first, "label"), ("Intentos de pago (#)", "int"), ("Pagos aprobados (#)", "int"),
                             ("Tasa de aprobación", "calc", 1, {"f": rate(2, 1), "total": True}),
                             ("Ventas aprobadas (COP)", "int"),
                             ("Participación en transacciones", "calc", 1, {"f": share(2)}),
                             ("Participación en ventas", "calc", 1, {"f": share(4)}),
                             ("Comentarios", "text", 1)]
    s.table("Volumen y aprobación por medio de pago", mixcols("Medio de pago"),
            [["PSE"], ["Tarjeta de crédito"], ["Tarjeta de débito"], ["Efectivo (indicar modalidad en comentarios)"],
             ["Saldo o crédito de Mercado Pago"], ["Otro (especificar en comentarios)"]],
            "A", "Ecommerce y Finanzas", "la base de todas las partes del caso", "Sí",
            dato=f"cerca de 60% PSE y 40% tarjetas ({D_SEP}); PSE con 58% de las ventas ({D_MAY})",
            example=["PSE", 1000, 900, 0.9, 45000000, 0.6, 0.6, "Texto libre"], total=True)
    s.table("Tarjetas por marca y tipo", mixcols("Marca y tipo"),
            [["Visa crédito"], ["Visa débito"], ["Mastercard crédito"], ["Mastercard débito"],
             ["American Express"], ["Diners Club"], ["Otra (especificar en comentarios)"]],
            "A", "Ecommerce y Finanzas", "aprobación por marca y costo transaccional", "Sí",
            example=["Visa crédito", 1000, 900, 0.9, 45000000, 0.3, 0.3, "Texto libre"], total=True)
    s.table("Tarjetas por origen y modalidad",
            [("Segmento", "label"), ("Intentos de pago (#)", "int"), ("Pagos aprobados (#)", "int"),
             ("Tasa de aprobación", "calc", 1, {"f": rate(2, 1)}),
             ("Ventas aprobadas (COP)", "int"), ("Comentarios", "text", 3)],
            [["Tarjetas emitidas en Colombia"], ["Tarjetas emitidas en el exterior"],
             ["Pagos con tarjeta guardada"], ["Pagos con tarjeta digitada en la compra"],
             ["Pagos a 1 cuota"], ["Pagos a más de 1 cuota"]],
            "B", "Ecommerce", "aprobación por segmento", "Parcial",
            note="Los segmentos no son excluyentes entre sí, por eso esta tabla no lleva total.",
            example=["Tarjetas emitidas en Colombia", 1000, 900, 0.9, 45000000, "Texto libre"])
    s.questions("Preguntas sobre efectivo y links de pago", [
        ("Modalidad de las transacciones en efectivo: pago contra entrega, recaudo en puntos físicos u otra. "
         "Indicar el proveedor si existe.", "A", "Ecommerce y Finanzas", "Base del caso", "No",
         f"Cerca de 20.000 transacciones al mes en efectivo ({D_MAY})."),
        ("Links de pago: para qué casos se usan y cuántos se generan y se pagan al mes.", "B", "Ecommerce",
         "Alcance de la integración", "Parcial", f"Usan links de pago de Mercado Pago para puntos de venta ({D_MAY})."),
    ])

    # ===== 4. Aprobación =====
    s = Sheet(wb, "4. Aprobación", 4, "4. Aprobación en tarjetas",
              f"Período: {PERIODO}, acumulado. Solo pagos con tarjeta de crédito y débito.")
    s.table("Aprobación por banco emisor: los 15 bancos con más intentos",
            [("Banco emisor", "label"), ("Tipo de tarjeta", "pick", 1, {"options": ["Crédito", "Débito", "Ambas"]}),
             ("Intentos de pago (#)", "int"), ("Pagos aprobados (#)", "int"),
             ("Tasa de aprobación", "calc", 1, {"f": rate(3, 2), "total": True}),
             ("Ventas aprobadas (COP)", "int"),
             ("Motivo de rechazo más frecuente", "text", 2)],
            [[None] for _ in range(15)], "A", "Ecommerce y Finanzas", "aprobación: dónde se concentran los rechazos", "Sí",
            dato=f"aprobación en tarjetas cercana a 90% ({D_SEP}); 91% sin reintentos ({D_MAY})",
            note="Si el reporte no trae el nombre del banco, el BIN (primeros 6 dígitos) nos sirve para identificarlo.",
            example=["Banco de ejemplo", "Crédito", 1000, 900, 0.9, 45000000, "Fondos insuficientes"],
            total=True)

    # Tabla de órdenes con filas de entrada y filas calculadas
    def ord_formula(kind):
        def f(cx):
            v = lambda i: f"{cx.c(1)}{cx.r1 + i}"
            if kind == "primer":
                n, d = v(1), v(0)
            elif kind == "recup":
                n, d = f"({v(3)}+{v(4)}+{v(5)})", v(2)
                return (f'=IF(AND(ISNUMBER({v(2)}),ISNUMBER({v(3)}),ISNUMBER({v(4)}),ISNUMBER({v(5)})),'
                        f'IF({d}>0,{n}/{d},""),"")')
            else:
                n, d = v(6), v(0)
            return f'=IF(AND(ISNUMBER({n}),ISNUMBER({d})),IF({d}>0,{n}/{d},""),"")'
        return f
    s.table("Aprobación por intento frente a aprobación por orden",
            [("Concepto", "label", 3), ("Valor", "int"), ("Comentarios", "text", 4)],
            [["Órdenes con pago por tarjeta iniciado (#)"],
             ["Órdenes aprobadas en el primer intento (#)"],
             ["Órdenes con el primer intento rechazado (#)"],
             ["   De esas, recuperadas con la misma tarjeta (#)"],
             ["   De esas, recuperadas con otra tarjeta (#)"],
             ["   De esas, recuperadas con otro medio de pago, por ejemplo PSE (#)"],
             ["   De esas, no recuperadas: venta perdida (#)"],
             ("CALC", "Calculado: tasa de aprobación en el primer intento", ord_formula("primer"), F_PCT),
             ("CALC", "Calculado: tasa de recuperación de las órdenes con primer intento rechazado",
              ord_formula("recup"), F_PCT),
             ("CALC", "Calculado: órdenes perdidas sobre órdenes iniciadas", ord_formula("perdida"), F_PCT)],
            "A", "Ecommerce", "aprobación: cuánta venta se pierde después de un rechazo", "Sí",
            example=["Órdenes con pago por tarjeta iniciado (#)", 1000, "Texto libre"])
    s.questions("Preguntas sobre aprobación", [
        ("¿La tasa de aprobación que manejan hoy se mide por intento o por orden?", "A", "Ecommerce", "Aprobación",
         "No", f"Cerca de 90% en tarjetas ({D_SEP})."),
        ("¿Hacen reintentos automáticos cuando un pago es rechazado? Si es así, con qué lógica y cuántas veces.",
         "A", "TI", "Aprobación", "No", f"91% de aprobación sin reintentos ({D_MAY})."),
        ("¿Qué ve el cliente cuando su pago es rechazado? Indicar si se le ofrece otro medio de pago.", "B",
         "Ecommerce", "Aprobación", "No", None),
        ("¿Mercado Pago aplica autenticación 3DS a sus pagos con tarjeta? Indicar a qué porcentaje de los pagos.",
         "B", "TI", "Aprobación", "No", None),
        ("Si aplica 3DS: porcentaje de autenticaciones exitosas, fallidas y abandonadas.", "C", "TI", "Aprobación",
         "Parcial", None),
        ("Patrones de rechazo por hora del día o día de la semana, si los han identificado.", "C", "Ecommerce",
         "Aprobación", "Sí", None),
    ])

    # ===== 5. Rechazos =====
    s = Sheet(wb, "5. Rechazos", 5, "5. Motivos de rechazo en tarjetas",
              f"Período: {PERIODO}, acumulado. Es la hoja más importante para la parte de aprobación del caso.")
    CLASES = ["Definitivo (hard)", "Temporal (soft)", "Riesgo", "Datos del comprador", "Otro"]
    SINCOD = "Sin código propio. Usar el código del banco si lo tienen."
    motivos = [
        ("Fondos o cupo insuficiente", "cc_rejected_insufficient_amount", CLASES[1]),
        ("Requiere autorización del banco", "cc_rejected_call_for_authorize", CLASES[1]),
        ("Rechazo genérico sin motivo específico", "cc_rejected_other_reason", CLASES[1]),
        ("Límite de intentos alcanzado", "cc_rejected_max_attempts", CLASES[1]),
        ("Error técnico o tiempo de espera agotado (banco, red o procesador)", SINCOD, CLASES[1]),
        ("Tarjeta inactiva o deshabilitada", "cc_rejected_card_disabled", CLASES[0]),
        ("Tarjeta vencida", SINCOD, CLASES[0]),
        ("Tarjeta reportada como robada o perdida", SINCOD, CLASES[0]),
        ("Cuenta cerrada o tarjeta cancelada", SINCOD, CLASES[0]),
        ("Tarjeta o comprador en lista de bloqueo", "cc_rejected_blacklist", CLASES[0]),
        ("Rechazo por prevención de fraude", "cc_rejected_high_risk", CLASES[2]),
        ("Número de tarjeta mal digitado", "cc_rejected_bad_filled_card_number", CLASES[3]),
        ("Fecha de vencimiento mal digitada", "cc_rejected_bad_filled_date", CLASES[3]),
        ("Código de seguridad mal digitado", "cc_rejected_bad_filled_security_code", CLASES[3]),
        ("Otros datos mal digitados", "cc_rejected_bad_filled_other", CLASES[3]),
        ("Número de cuotas no válido", "cc_rejected_invalid_installments", CLASES[3]),
        ("Pago duplicado", "cc_rejected_duplicated_payment", CLASES[4]),
    ]
    rows = [[m, c, k] for m, c, k in motivos] + [[None, None, None] for _ in range(4)]
    tid, r1, r2, lt = s.table(
        "Rechazos por motivo",
        [("Motivo del rechazo", "label"), ("Código de referencia en Mercado Pago (status_detail)", "text", 2),
         ("Clasificación", "pick", 1, {"options": CLASES}),
         ("Rechazos (#)", "int"), ("Participación en rechazos", "calc", 1, {"f": share(3)}),
         ("Monto rechazado (COP)", "int"), ("Comentarios", "text", 1)],
        rows, "A", "Ecommerce y TI", "aprobación: qué parte de los rechazos es recuperable", "Sí",
        dato=f"la mayoría de los rechazos son declinaciones definitivas del banco emisor ({D_SEP})",
        note="La clasificación es la propuesta de trabajo de Yuno para este análisis y se puede cambiar con la lista de "
             "cada celda. Los códigos son los que documenta Mercado Pago en el campo status_detail. Las cuatro filas "
             "en blanco son para motivos que no estén en la lista.",
        example=["Fondos o cupo insuficiente", "cc_rejected_insufficient_amount", CLASES[1], 300, 0.3, 13500000,
                 "Texto libre"], total=True)
    # Los códigos y la clasificación vienen diligenciados: se muestran como referencia, no como celdas por llenar
    for i in range(len(motivos)):
        for col in (3, 4):
            c = s.ws.cell(row=r1 + i, column=col)
            c.fill, c.font = fill(WHITE), font(9, color=DGRAY)
        s.ws.cell(row=r1 + i, column=5).fill = fill(WHITE)

    cls_rng = f"${lt[2]}${r1}:${lt[2]}${r2}"
    num_rng = f"${lt[3]}${r1}:${lt[3]}${r2}"

    def sumif(cx):
        return f'=SUMIF({cls_rng},{cx.c(0)}{cx.r},{num_rng})'
    s.table("Separación entre rechazos definitivos y temporales",
            [("Tipo de rechazo", "label"), ("Qué significa", "fixed", 2),
             ("Según Cruz Verde: rechazos (#)", "int"),
             ("Según Cruz Verde: participación", "calc", 1, {"f": share(2)}),
             ("Calculado desde la tabla 5.1 (#)", "calc", 1, {"f": sumif, "fmt": F_INT}),
             ("Calculado: participación", "calc", 1, {"f": share(4)}),
             ("Comentarios", "text", 1)],
            [[CLASES[0], "El banco indica que esa tarjeta no puede aprobar. Reintentar con la misma tarjeta no cambia el resultado."],
             [CLASES[1], "Depende del momento o de la ruta. Puede aprobarse con un reintento, por otro procesador o más tarde."],
             [CLASES[2], "El pago lo detuvo el motor de prevención de fraude del proveedor o del banco."],
             [CLASES[3], "El cliente digitó mal un dato. Se corrige en el checkout."],
             [CLASES[4], "Motivos que no entran en las categorías anteriores."]],
            "A", "Ecommerce y TI", "aprobación: el potencial real de recuperación", "Sí",
            dato=f"la mayoría de los rechazos son declinaciones definitivas del banco emisor ({D_SEP})",
            note="Si ya tienen su propia separación, escríbanla en la columna \"Según Cruz Verde\". Las columnas "
                 "calculadas se llenan solas con la tabla 5.1 y sirven para contrastar las dos lecturas.",
            total=True)
    s.questions("Preguntas sobre rechazos", [
        ("¿De dónde sale hoy la clasificación de un rechazo como definitivo? Por ejemplo: reporte de Mercado Pago, "
         "análisis interno, información del banco.", "A", "Ecommerce", "Aprobación", "No",
         f"La mayoría de los rechazos son declinaciones definitivas del banco emisor ({D_SEP})."),
        ("¿Tienen acceso al código de respuesta original del banco emisor o solo al motivo que entrega Mercado Pago?",
         "A", "TI", "Aprobación", "No", None),
        ("Rechazos por prevención de fraude: ¿pueden ajustar las reglas o es una decisión cerrada del proveedor?",
         "B", "TI", "Aprobación", "No", f"Fraude por debajo de 0,05% ({D_MAY})."),
        ("¿Los rechazos se concentran en algún rango de ticket, canal o tipo de tarjeta?", "B", "Ecommerce",
         "Aprobación", "Sí", None),
        ("¿Han medido cuántos clientes con pago rechazado terminan comprando en tienda física o no compran?", "C",
         "Ecommerce", "Aprobación", "No", None),
    ])

    # ===== 6. PSE y efectivo =====
    s = Sheet(wb, "6. PSE y efectivo", 6, "6. PSE y efectivo", f"Período: {PERIODO}, acumulado. Cifras en COP.")
    s.table("Embudo de pagos por PSE",
            [("Etapa", "label", 3), ("Transacciones (#)", "int"),
             ("Porcentaje sobre iniciados", "calc", 1, {"f": over_first(1)}),
             ("Monto (COP)", "int"), ("Comentarios", "text", 2)],
            [["Pagos PSE iniciados: el cliente eligió PSE y fue enviado a su banco"],
             ["Aprobados"], ["Rechazados por el banco"], ["Pendientes que nunca se confirmaron"],
             ["Abandonados o expirados"], ["Fallidos por error técnico o caída del servicio"]],
            "A", "Ecommerce y TI", "aprobación real de PSE y continuidad", "Sí",
            dato=f"aprobación por PSE cercana a 100% ({D_SEP}). Queremos confirmar si se mide sobre pagos iniciados o sobre pagos finalizados",
            example=["Aprobados", 900, 0.9, 45000000, "Texto libre"])
    s.table("PSE por banco: los 10 bancos con más pagos iniciados",
            [("Banco", "label"), ("Pagos iniciados (#)", "int"), ("Pagos aprobados (#)", "int"),
             ("Tasa de aprobación", "calc", 1, {"f": rate(2, 1), "total": True}),
             ("Monto aprobado (COP)", "int"), ("Comentarios", "text", 3)],
            [[None] for _ in range(10)], "B", "Ecommerce y TI", "aprobación de PSE por banco", "Sí",
            example=["Banco de ejemplo", 1000, 900, 0.9, 45000000, "Texto libre"], total=True)
    s.table("Embudo de pagos en efectivo",
            [("Etapa", "label", 3), ("Órdenes (#)", "int"),
             ("Porcentaje sobre generadas", "calc", 1, {"f": over_first(1)}),
             ("Monto (COP)", "int"), ("Comentarios", "text", 2)],
            [["Órdenes generadas con pago en efectivo"], ["Pagadas"], ["No pagadas o canceladas"]],
            "B", "Ecommerce y Finanzas", "la base del caso y nuevos medios de pago", "Parcial",
            dato=f"cerca de 20.000 transacciones al mes en efectivo ({D_MAY})",
            example=["Pagadas", 800, 0.8, 24000000, "Texto libre"])
    s.questions("Preguntas sobre PSE", [
        ("Tiempo promedio entre el inicio del pago por PSE y su confirmación.", "B", "TI", "Continuidad", "Sí", None),
        ("Pagos por PSE confirmados después de que la orden se canceló o expiró: cuántos al mes y cómo los manejan.",
         "B", "Ecommerce y Finanzas", "Devoluciones y conciliación", "Parcial", None),
    ])

    # ===== 7. Incidentes =====
    s = Sheet(wb, "7. Incidentes", 7, "7. Incidentes y disponibilidad", f"Período: {PERIODO}.")
    s.table("Registro de caídas o intermitencias",
            [("Medio o servicio afectado", "label"), ("Fecha", "text"), ("Duración (minutos)", "int"),
             ("Alcance", "pick", 1, {"options": ["Total", "Parcial"]}),
             ("Transacciones afectadas, estimado (#)", "int"), ("Ventas afectadas, estimado (COP)", "int"),
             ("Cómo se detectó", "text", 1), ("Acción tomada", "text", 1)],
            [[None] for _ in range(10)], "A", "TI y Ecommerce", "continuidad: el costo de depender de una sola ruta", "No",
            dato=f"caídas de PSE y dependencia de un solo proveedor entre los dolores mencionados ({D_MAY})",
            example=["PSE", "2026-03-14", 90, "Total", 300, 15000000, "Reclamos de clientes",
                     "Se esperó a que el servicio volviera"], total=False)
    s.questions("Preguntas sobre disponibilidad", [
        ("Número de caídas o intermitencias de PSE en los últimos 12 meses.", "A", "TI", "Continuidad", "No",
         f"Caídas de PSE mencionadas como dolor ({D_MAY})."),
        ("Tiempo total sin servicio de pagos en los últimos 12 meses (horas).", "A", "TI", "Continuidad", "No", None),
        ("¿Qué alternativa tiene el cliente cuando PSE o Mercado Pago no responde?", "A", "Ecommerce y TI",
         "Continuidad", "No", f"Mercado Pago es el único proveedor ({D_SEP})."),
        ("Ventas promedio por hora en horario de mayor demanda y en horario normal (COP).", "B", "Ecommerce",
         "Continuidad", "Sí", None),
        ("¿Cómo monitorean hoy la disponibilidad de los pagos? Por ejemplo: alertas, tableros, avisos del proveedor.",
         "B", "TI", "Continuidad", "No", None),
    ])

    # ===== 8. Devoluciones =====
    s = Sheet(wb, "8. Devoluciones", 8, "8. Devoluciones, contracargos y fraude",
              f"Período: {PERIODO}, acumulado. Cifras en COP.")
    s.table("Devoluciones por medio de pago",
            [("Medio de pago", "label"), ("Devoluciones (#)", "int"), ("Monto devuelto (COP)", "int"),
             ("Devoluciones totales (#)", "int"), ("Devoluciones parciales (#)", "int"),
             ("Días hasta que el cliente recibe el dinero", "int", 1, {"nototal": True}),
             ("Costo por devolución (COP)", "int", 1, {"nototal": True}),
             ("Comentarios", "text", 1)],
            [["PSE"], ["Tarjeta de crédito"], ["Tarjeta de débito"], ["Efectivo"], ["Otro (especificar)"]],
            "A", "Finanzas y Ecommerce", "devoluciones: volumen, tiempos y costo actual", "Parcial",
            dato=f"devoluciones cercanas a 5% de las compras, concentradas en PSE ({D_SEP})",
            note="Las columnas de días y de costo son promedios por medio de pago, por eso no llevan total.",
            example=["PSE", 50, 2500000, 40, 10, 8, 3000, "Texto libre"], total=True)
    s.table("Motivos de devolución",
            [("Motivo", "label", 3), ("Devoluciones (#)", "int"),
             ("Participación", "calc", 1, {"f": share(1)}), ("Monto devuelto (COP)", "int"),
             ("Comentarios", "text", 2)],
            [["Producto agotado o no disponible"], ["Cancelación solicitada por el cliente"],
             ["Pedido no entregado o fuera de cobertura"], ["Error de precio o promoción"], ["Cobro duplicado"],
             ["Pago confirmado después de cancelar la orden"], [None], [None]],
            "A", "Ecommerce", "devoluciones: cuáles se pueden evitar y cuáles resolver con saldo", "No",
            note="Las dos filas en blanco son para motivos propios de Cruz Verde que no estén en la lista.",
            example=["Producto agotado o no disponible", 50, 0.5, 2500000, "Texto libre"], total=True)
    s.questions("Preguntas sobre devoluciones", [
        ("Describir el proceso actual de una devolución por PSE, paso a paso: quién la solicita, quién la aprueba "
         "y cómo se gira el dinero.", "A", "Finanzas", "Devoluciones", "No",
         f"Devoluciones de PSE mencionadas como dolor ({D_MAY})."),
        ("¿La devolución por PSE la ejecuta Mercado Pago o la hace Cruz Verde por transferencia bancaria?", "A",
         "Finanzas", "Devoluciones", "No", None),
        ("Solicitudes o reclamos de clientes relacionados con devoluciones, promedio al mes.", "B", "Ecommerce",
         "Devoluciones", "No", None),
        ("Porcentaje de las devoluciones que podría entregarse como saldo o crédito Cruz Verde si existiera esa opción.",
         "B", "Ecommerce y Finanzas", "Devoluciones", "No",
         f"Interés en gift cards o créditos Cruz Verde para devoluciones y pagos ({D_SEP})."),
    ])
    s.questions("Preguntas sobre contracargos y fraude", [
        ("Contracargos en los últimos 12 meses: número y monto (COP).", "B", "Finanzas", "Costo transaccional",
         "Parcial", f"Contracargos muy bajos ({D_SEP})."),
        ("Tasa de contracargos sobre los pagos aprobados con tarjeta (%).", "B", "Finanzas", "Costo transaccional",
         "Parcial", None),
        ("Motivos principales de contracargo.", "C", "Finanzas", "Costo transaccional", "No", None),
        ("Porcentaje de contracargos ganados en disputa.", "C", "Finanzas", "Costo transaccional", "No", None),
        ("Tasa de fraude confirmado sobre los pagos aprobados (%).", "B", "Finanzas", "Aprobación", "No",
         f"Por debajo de 0,05% ({D_MAY})."),
        ("Herramientas o reglas de prevención de fraude que usan además de las de Mercado Pago.", "C", "TI",
         "Aprobación", "No", f"Una herramienta antifraude no es prioridad en esta etapa ({D_SEP})."),
    ])

    # ===== 9. Costos =====
    s = Sheet(wb, "9. Costos", 9, "9. Costos actuales de pagos",
              "Tarifas vigentes y costo promedio de los últimos 3 meses. Cifras en COP.")
    s.table("Tarifas por medio de pago",
            [("Medio de pago", "label"), ("Comisión variable (%)", "pct"), ("Comisión fija por transacción (COP)", "int"),
             ("¿La tarifa incluye IVA?", "yn"), ("Plazo de liquidación (días)", "int"),
             ("Costo mensual promedio (COP)", "int"), ("Comentarios", "text", 2)],
            [["PSE"], ["Tarjeta de crédito"], ["Tarjeta de débito"], ["Efectivo"], ["Link de pago"],
             ["Otro (especificar)"]],
            "A", "Finanzas", "costo transaccional: la línea base contra la que se mide el ahorro", "Parcial",
            example=["PSE", 0.025, 900, "No", 2, 12000000, "Texto libre"])
    s.table("Otros costos asociados a pagos",
            [("Concepto", "label"), ("Costo (COP)", "int"),
             ("Unidad", "pick", 1, {"options": ["Por evento", "Mensual", "Anual", "Porcentaje"]}),
             ("Costo mensual promedio (COP)", "int"), ("Comentarios", "text", 4)],
            [["Devolución"], ["Contracargo"], ["Retiro o transferencia a cuenta bancaria"],
             ["Adelanto de liquidación"], ["Prevención de fraude"], ["Autenticación 3DS"],
             ["Cargo fijo o mensualidad de plataforma"], ["Otro (especificar)"]],
            "B", "Finanzas", "costo transaccional: el costo total de la operación de pagos", "No",
            example=["Devolución", 3000, "Por evento", 150000, "Texto libre"])
    s.questions("Preguntas sobre costos y contratos", [
        ("Costo total mensual de procesamiento de pagos del canal digital, promedio de los últimos 3 meses (COP).",
         "A", "Finanzas", "Costo transaccional", "Parcial", None),
        ("Documento de soporte: últimas 3 liquidaciones o facturas mensuales de Mercado Pago.", "A", "Finanzas",
         "Costo transaccional", "No", None),
        ("Vigencia del contrato con Mercado Pago y condiciones relevantes: exclusividad, volúmenes mínimos, "
         "descuentos por volumen.", "A", "Finanzas", "Costo transaccional", "No",
         f"Mercado Pago es el único proveedor ({D_SEP})."),
        ("Retenciones que se aplican en la liquidación (retención en la fuente, ReteICA, ReteIVA) y cómo las "
         "registran.", "B", "Finanzas", "Conciliación", "No", None),
        ("Tarifas ofrecidas por los proveedores que han evaluado, si pueden compartirlas.", "C", "Finanzas",
         "Costo transaccional", "No", f"Proveedores conversados: PayU, Redeban, Movii y Kushki ({D_SEP})."),
        ("Moneda y periodicidad en que prefieren recibir la facturación de Yuno.", "B", "Finanzas",
         "Propuesta comercial", "No", f"Prefieren negociar en pesos ({D_SEP})."),
    ])

    # ===== 10. Conciliación =====
    s = Sheet(wb, "10. Conciliación", 10, "10. Conciliación",
              "Proceso actual de conciliación de los pagos del canal digital.")
    s.questions("Preguntas sobre el proceso de conciliación", [
        ("Describir el proceso actual de conciliación de pagos digitales, paso a paso.", "A", "Finanzas",
         "Conciliación", "No", f"La conciliación es el primer frente de interés ({D_SEP})."),
        ("Fuentes que cruzan hoy. Por ejemplo: reporte de Mercado Pago, extractos bancarios, sistema de pedidos, ERP.",
         "A", "Finanzas", "Conciliación", "No", None),
        ("Sistema contable o ERP donde se registra el recaudo.", "A", "Finanzas y TI", "Conciliación", "No", None),
        ("Frecuencia de la conciliación: diaria, semanal o mensual.", "A", "Finanzas", "Conciliación", "No", None),
        ("Número o porcentaje de partidas que no cruzan de forma automática cada mes.", "A", "Finanzas",
         "Conciliación", "No", None),
        ("Causas más frecuentes de diferencias. Por ejemplo: devoluciones, retenciones, comisiones, pagos "
         "duplicados, pagos de PSE confirmados tarde.", "A", "Finanzas", "Conciliación", "No", None),
        ("Días que toma el cierre mensual de pagos.", "B", "Finanzas", "Conciliación", "No", None),
        ("Monto promedio de partidas pendientes por conciliar al cierre de mes (COP).", "B", "Finanzas",
         "Conciliación", "No", None),
        ("Número de cuentas bancarias donde reciben el recaudo digital.", "B", "Finanzas", "Conciliación", "No", None),
        ("Nivel de detalle que necesitan: por transacción, por liquidación, por tienda o por canal.", "B", "Finanzas",
         "Conciliación", "No", None),
        ("Formato de archivo que requiere su sistema contable para cargar el recaudo.", "B", "Finanzas y TI",
         "Conciliación", "No", None),
        ("Documento de soporte: un ejemplo del reporte de liquidación de Mercado Pago que usan hoy. Puede ser de "
         "un solo día.", "B", "Finanzas", "Conciliación", "No", None),
        ("Horas al mes dedicadas a la conciliación de pagos digitales. Opcional: solo sirve para dimensionar el "
         "alcance y no entra a la sección de ahorros del business case.", "C", "Finanzas", "Conciliación", "No",
         None),
    ])

    # ===== 11. Checkout =====
    s = Sheet(wb, "11. Checkout", 11, "11. Checkout y tecnología",
              "Embudo: promedio mensual de los últimos 3 meses.")
    s.table("Embudo de checkout",
            [("Etapa", "label"), ("Sitio web: sesiones (#)", "int"), ("App: sesiones (#)", "int"),
             ("Sitio web: paso desde la etapa anterior", "calc", 1, {"f": over_prev(1)}),
             ("App: paso desde la etapa anterior", "calc", 1, {"f": over_prev(2)}),
             ("Comentarios", "text", 3)],
            [["Carrito con productos"], ["Inicio del checkout"], ["Selección del medio de pago"],
             ["Pago enviado"], ["Pago aprobado"]],
            "B", "Ecommerce", "nuevos medios de pago: dónde se pierde el cliente antes de pagar", "No",
            example=["Inicio del checkout", 10000, 8000, 0.5, 0.5, "Texto libre"])
    s.questions("Preguntas sobre checkout y tecnología", [
        ("Plataforma de ecommerce: desarrollo propio u otra (indicar cuál).", "A", "TI", "Alcance de la integración",
         "No", None),
        ("Tipo de integración con Mercado Pago: checkout con redirección, checkout integrado por API o links de pago.",
         "A", "TI", "Alcance de la integración", "No", f"Mercado Pago único, incluidos links de pago ({D_MAY})."),
        ("¿El cliente sale del sitio o de la app para completar el pago?", "B", "TI", "Aprobación", "No", None),
        ("Participación de las ventas por dispositivo: iOS, Android y computador (%).", "A", "Ecommerce",
         "Nuevos medios de pago", "Parcial", None),
        ("¿Permiten guardar tarjetas? Indicar cuántas tarjetas guardadas activas tienen.", "B", "TI",
         "Aprobación", "No", None),
        ("Sistemas que consumen la información de pagos: sistema de pedidos, ERP, facturación electrónica.", "B",
         "TI", "Alcance de la integración", "No", None),
        ("Restricciones o ventanas de congelamiento de TI en los próximos 6 meses.", "B", "TI",
         "Plan de implementación", "No", None),
    ])

    # ===== 12. Nuevos medios =====
    s = Sheet(wb, "12. Nuevos medios", 12, "12. Nuevos medios de pago y planes",
              "Prioridades de Cruz Verde y evidencia de demanda.")
    s.table("Medios de pago de interés",
            [("Medio de pago", "label"),
             ("Prioridad para Cruz Verde", "pick", 1, {"options": ["1 Alta", "2 Media", "3 Baja", "No interesa"]}),
             ("¿Lo aceptan hoy en tiendas físicas?", "yn"),
             ("Participación en ventas de tienda física (%)", "pct"),
             ("Evidencia de demanda en el canal digital", "text", 2),
             ("Proveedor con el que lo han conversado", "text", 1),
             ("Fecha deseada de salida", "text", 1)],
            [["Nequi"], ["Daviplata"], ["Bre-B"], ["Apple Pay"], ["Google Pay"], ["Addi"], ["Sistecrédito"],
             ["Gift cards o créditos Cruz Verde"], [None], [None]],
            "A", "Ecommerce", "nuevos medios de pago: el volumen adicional de cada uno", "No",
            dato=f"interés en Sistecrédito y en gift cards o créditos Cruz Verde ({D_SEP})",
            note="Las dos filas en blanco son para otros medios de pago que tengan en el radar.",
            example=["Nequi", "1 Alta", "Sí", 0.08, "Solicitudes en servicio al cliente", "Por definir",
                     "Primer trimestre de 2027"])
    s.questions("Preguntas sobre demanda y planes", [
        ("Mix de medios de pago en tiendas físicas: efectivo, débito, crédito, billeteras y otros (%). "
         "Nos sirve como referencia de la preferencia del cliente.", "B", "Finanzas", "Nuevos medios de pago", "No",
         None),
        ("Solicitudes de clientes por medios de pago que hoy no ofrecen: encuestas o casos de servicio al cliente.",
         "B", "Ecommerce", "Nuevos medios de pago", "No", None),
        ("Estado de las conversaciones con PayU, Redeban, Movii y Kushki: propuesta recibida, contrato o integración.",
         "A", "Ecommerce y Finanzas", "Alcance de las conexiones", "No",
         f"Proveedores conversados: PayU, Redeban, Movii y Kushki ({D_SEP})."),
        ("Porcentaje estimado de clientes del canal digital que no tienen tarjeta de crédito.", "C", "Ecommerce",
         "Nuevos medios de pago", "No", None),
        ("Clientes con compras recurrentes del mismo producto: cuántos son y con qué frecuencia compran.", "C",
         "Ecommerce", "Suscripciones a futuro", "Parcial", f"Motor de suscripciones como interés a futuro ({D_SEP})."),
        ("Otros proyectos de pagos previstos para 2027.", "C", "Ecommerce y TI", "Plan de implementación", "No", None),
    ])

    # ===== Resumen =====
    ws = ws_res
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = BLACK
    widths = [7, 19, 62, 11, 22, 30, 16, 46, 16, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
    nc = len(widths)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nc)
    c = ws.cell(row=1, column=1, value="Resumen de la información solicitada")
    c.font, c.alignment = font(16, True, color=WHITE), Alignment(vertical="center", indent=1)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=nc)
    c2 = ws.cell(row=2, column=1, value="Lista completa de solicitudes. Usen las columnas Estado y Comentarios para "
                                         "llevar el avance. El nombre de la hoja lleva a la tabla correspondiente.")
    c2.font, c2.alignment = font(10, color=DGRAY), Alignment(vertical="center", indent=1)
    for col in range(1, nc + 1):
        ws.cell(row=1, column=col).fill = fill(BLUE)
        ws.cell(row=2, column=col).fill = fill(LILAC)
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 34, 24

    HR = 6  # fila de encabezado de la lista
    first, last = HR + 1, HR + len(RESUMEN)
    cnt = lambda col, val: f'COUNTIF(${col}${first}:${col}${last},"{val}")'
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=nc)
    st = ws.cell(row=4, column=1, value=(
        f'="Solicitudes: "&COUNTA($A${first}:$A${last})&".   Prioridad A: "&{cnt("D", "A")}'
        f'&".   Prioridad B: "&{cnt("D", "B")}&".   Prioridad C: "&{cnt("D", "C")}'
        f'&".   Enviadas: "&{cnt("I", "Enviado")}&".   Pendientes: "&{cnt("I", "Pendiente")}&"."'))
    st.font, st.alignment = font(10, True), Alignment(vertical="center", indent=1)
    for col in range(1, nc + 1):
        ws.cell(row=4, column=col).fill = fill(CALC)
    ws.row_dimensions[4].height = 24

    heads = ["ID", "Hoja", "Información solicitada", "Prioridad", "Área sugerida", "Para qué la usamos",
             "Se calcula con el export", "Dato que tenemos hoy, a confirmar", "Estado", "Comentarios de Cruz Verde"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(row=HR, column=j, value=h)
        c.font, c.fill, c.border = font(10, True, color=WHITE), fill(BLUE), BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[HR].height = 32
    dv = DataValidation(type="list", formula1='"Pendiente,Enviado,Parcial,No disponible"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, (rid, hoja, texto, prio, area, uso, exp, dato) in enumerate(RESUMEN):
        r = first + i
        uso = uso[0].upper() + uso[1:]
        if dato:
            dato = dato[0].upper() + dato[1:]
            dato = dato if dato.endswith(".") else dato + "."
        vals = [rid, hoja, texto, prio, area, uso, exp, dato, "Pendiente", None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = BORDER
            c.font = font(9 if j in (1, 8) else 10, bold=(j == 4), italic=(j == 8),
                          color=DGRAY if j in (1, 8) else BLACK)
            c.fill = fill(INPUT if j in (9, 10) else WHITE)
            c.alignment = Alignment(wrap_text=True, vertical="top",
                                    horizontal="center" if j in (1, 4, 7, 9) else "left")
        link = ws.cell(row=r, column=2)
        link.hyperlink = f"#'{hoja}'!A1"
        link.font = font(10, color=BLUE)
        dv.add(f"I{r}")
        ws.row_dimensions[r].height = row_height(max(
            n_lines(texto, widths[2] - 1), n_lines(dato, widths[7] - 1, 9), n_lines(uso, widths[5] - 1),
            n_lines(area, widths[4] - 1)))
    ws.freeze_panes = f"D{first}"
    ws.auto_filter.ref = f"A{HR}:{L(nc)}{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.active = 0
    wb.save(OUT)
    return OUT, len(RESUMEN)


if __name__ == "__main__":
    path, n = build()
    print(f"OK {path} ({n} solicitudes)")
